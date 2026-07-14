"""Unified, fail-closed tool execution surface for Kolibri.

Every implemented handler returns either a concrete result or a verified
artifact.  Unsupported or unproved tools are rejected with a structured
capability error; there is no chat fallback that can pretend a file, image or
preview was created.
"""

from __future__ import annotations

import base64
from io import BytesIO
import hashlib
import hmac
import json
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile
import time
from typing import Any
from urllib.parse import urlparse

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field, ValidationError

from app.artifact_store import (
    ArtifactIntegrityError,
    ArtifactNotFound,
    ArtifactStore,
    ArtifactStoreError,
    ArtifactValidationError,
    get_artifact_store,
)
from app.browser_session import ProjectPrincipal, resolve_project_principal
from app.capability_runtime import (
    capability_by_id,
    record_capability_invocation,
)
from app.routers.openai_compat import _authorize_api_key_admin


router = APIRouter(tags=["tools"])
_TOOL_ID = re.compile(r"^[a-z][a-z0-9_.-]{1,79}$")
_MAX_UPLOAD_BYTES = 100 * 1024 * 1024
_MAX_ANALYSIS_TEXT = 2_000_000


class ToolInvocationRequest(BaseModel):
    tool: str = Field(min_length=2, max_length=80)
    arguments: dict[str, Any] = Field(default_factory=dict)


class CapabilityInvocationError(RuntimeError):
    def __init__(self, capability_id: str, reason_code: str, *, retryable: bool = False):
        self.capability_id = capability_id
        self.reason_code = reason_code
        self.retryable = retryable
        super().__init__(f"{capability_id}: {reason_code}")


def _capability_error(exc: CapabilityInvocationError) -> HTTPException:
    return HTTPException(
        status_code=503,
        detail={
            "code": "capability_unavailable",
            "message": "Запрошенная возможность сейчас недоступна.",
            "capability": exc.capability_id,
            "reason_code": exc.reason_code,
            "retryable": exc.retryable,
        },
    )


def _assert_probeable(capability_id: str) -> dict[str, Any]:
    """Allow a configured degraded route to perform its first real probe."""

    capability = capability_by_id(capability_id)
    if capability is None:
        raise CapabilityInvocationError(capability_id, "capability_not_registered")
    if capability.get("status") == "available" and capability.get("invocable") is True:
        return capability

    reason = capability.get("reason") if isinstance(capability.get("reason"), dict) else {}
    reason_code = str(reason.get("code") or "capability_unavailable")
    policy = capability.get("policy") if isinstance(capability.get("policy"), dict) else {}
    renderer = capability.get("renderer") if isinstance(capability.get("renderer"), dict) else {}
    routes = capability.get("routes") if isinstance(capability.get("routes"), list) else []
    eligible_route = any(
        isinstance(route, dict)
        and route.get("configured") is True
        and route.get("permitted") is True
        and isinstance(route.get("credential"), dict)
        and route["credential"].get("ready") is True
        for route in routes
    )
    renderer_ready = not renderer.get("required") or renderer.get("registered") is True
    if policy.get("permitted") is not True or not renderer_ready or not eligible_route:
        raise CapabilityInvocationError(capability_id, reason_code)
    if reason_code not in {
        "probe_not_run",
        "probe_stale",
        "renderer_health_unverified",
        "invocation_failed",
    }:
        raise CapabilityInvocationError(capability_id, reason_code)
    return capability


def _record_success(capability_id: str, artifact: dict[str, Any] | None = None) -> None:
    evidence_id = str(artifact.get("sha256") or "") if isinstance(artifact, dict) else None
    record_capability_invocation(
        capability_id,
        succeeded=True,
        provider="kolibri-backend",
        evidence_id=evidence_id or None,
    )


def _record_failure(capability_id: str, code: str) -> None:
    record_capability_invocation(
        capability_id,
        succeeded=False,
        error_code=code,
        provider="kolibri-backend",
    )


def _bounded_text(value: Any, *, field: str, maximum: int = 1_000_000) -> str:
    if not isinstance(value, str):
        raise HTTPException(status_code=422, detail={"code": f"{field}_must_be_text"})
    text = value.strip()
    if not text or len(text) > maximum:
        raise HTTPException(status_code=422, detail={"code": f"{field}_invalid"})
    return text


def _safe_filename(value: Any, default: str) -> str:
    filename = str(value or default).strip()
    if filename != filename.rsplit("/", 1)[-1] or "\\" in filename or "\x00" in filename:
        raise HTTPException(status_code=422, detail={"code": "filename_invalid"})
    return filename[:240]


def _store_artifact(
    content: bytes,
    *,
    capability_id: str,
    artifact_type: str,
    mime_type: str,
    filename: str,
    title: str,
    principal: ProjectPrincipal,
    metadata: dict[str, Any] | None = None,
) -> dict[str, Any]:
    merged_metadata = {
        # Never persist or expose the raw anonymous session identifier.  The
        # hash is sufficient for exact tenant filtering and cannot be used as
        # the signed browser cookie.
        "scope_key": hashlib.sha256(principal.scope_id.encode()).hexdigest(),
        "producer_capability": capability_id,
        **(metadata or {}),
    }
    try:
        artifact = get_artifact_store().put_bytes(
            content,
            artifact_type=artifact_type,
            mime_type=mime_type,
            filename=filename,
            title=title,
            metadata=merged_metadata,
        )
        # Re-open before reporting success.  This proves persistence, MIME,
        # size and digest rather than trusting the generator return value.
        verified = get_artifact_store().open(artifact["id"])
    except ArtifactStoreError as exc:
        _record_failure(capability_id, type(exc).__name__.lower())
        raise HTTPException(
            status_code=502,
            detail={"code": "artifact_verification_failed", "capability": capability_id},
        ) from exc
    if verified.manifest != artifact or not verified.content:
        _record_failure(capability_id, "artifact_reopen_mismatch")
        raise HTTPException(status_code=502, detail={"code": "artifact_reopen_mismatch"})
    _record_success(capability_id, artifact)
    return artifact


async def _web_search(arguments: dict[str, Any]) -> dict[str, Any]:
    from app.web_search import web_search

    query = _bounded_text(arguments.get("query"), field="query", maximum=2_000)
    limit = arguments.get("limit", 5)
    if not isinstance(limit, int) or not 1 <= limit <= 20:
        raise HTTPException(status_code=422, detail={"code": "limit_invalid"})
    results = await web_search(query, limit)
    sources: list[dict[str, str]] = []
    for item in results:
        if not isinstance(item, dict):
            continue
        url = str(item.get("url") or "").strip()
        parsed = urlparse(url)
        if parsed.scheme not in {"http", "https"} or not parsed.netloc:
            continue
        sources.append(
            {
                "title": str(item.get("title") or url)[:500],
                "snippet": str(item.get("snippet") or "")[:4_000],
                "url": url,
            }
        )
    if not sources:
        _record_failure("web.search", "no_verified_sources")
        raise HTTPException(
            status_code=502,
            detail={"code": "web_search_no_verified_sources", "retryable": True},
        )
    record_capability_invocation(
        "web.search",
        succeeded=True,
        provider="web-search",
        evidence_id="sha256:" + json_hash(sources),
    )
    return {"query": query, "sources": sources, "total": len(sources)}


def json_hash(value: Any) -> str:
    encoded = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    return hashlib.sha256(encoded).hexdigest()


def _generate_pdf(arguments: dict[str, Any], principal: ProjectPrincipal) -> dict[str, Any]:
    from app.pdf_generator import generate_document_pdf

    title = _bounded_text(arguments.get("title"), field="title", maximum=240)
    content = _bounded_text(arguments.get("content"), field="content")
    _assert_probeable("document.pdf")
    try:
        data = generate_document_pdf({"title": title, "content": content})
    except Exception as exc:
        _record_failure("document.pdf", "generator_failed")
        raise HTTPException(status_code=502, detail={"code": "pdf_generation_failed"}) from exc
    filename = _safe_filename(arguments.get("filename"), f"{title}.pdf")
    if not filename.lower().endswith(".pdf"):
        filename += ".pdf"
    return _store_artifact(
        data,
        capability_id="document.pdf",
        artifact_type="document.pdf",
        mime_type="application/pdf",
        filename=filename,
        title=title,
        principal=principal,
    )


def _generate_docx(arguments: dict[str, Any], principal: ProjectPrincipal) -> dict[str, Any]:
    from app.docx_export import html_to_docx

    title = _bounded_text(arguments.get("title"), field="title", maximum=240)
    content = _bounded_text(arguments.get("content"), field="content")
    _assert_probeable("document.docx")
    try:
        data = html_to_docx(title, content)
    except Exception as exc:
        _record_failure("document.docx", "generator_failed")
        raise HTTPException(status_code=502, detail={"code": "docx_generation_failed"}) from exc
    filename = _safe_filename(arguments.get("filename"), f"{title}.docx")
    if not filename.lower().endswith(".docx"):
        filename += ".docx"
    return _store_artifact(
        data,
        capability_id="document.docx",
        artifact_type="document.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=filename,
        title=title,
        principal=principal,
    )


def _generate_xlsx(arguments: dict[str, Any], principal: ProjectPrincipal) -> dict[str, Any]:
    from app.xlsx_export import generate_estimate_xlsx

    estimate = arguments.get("estimate")
    if not isinstance(estimate, dict):
        raise HTTPException(status_code=422, detail={"code": "estimate_must_be_object"})
    title = _bounded_text(estimate.get("title") or arguments.get("title"), field="title", maximum=240)
    _assert_probeable("document.xlsx")
    try:
        data = generate_estimate_xlsx(estimate)
    except Exception as exc:
        # This endpoint classifies the exception as invalid source data (422),
        # not a runtime outage.  A bad client payload must not poison the
        # capability's last successful execution proof.
        raise HTTPException(status_code=422, detail={"code": "xlsx_source_invalid"}) from exc
    filename = _safe_filename(arguments.get("filename"), f"{title}.xlsx")
    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"
    return _store_artifact(
        data,
        capability_id="document.xlsx",
        artifact_type="document.xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=filename,
        title=title,
        principal=principal,
    )


def _generate_pptx(arguments: dict[str, Any], principal: ProjectPrincipal) -> dict[str, Any]:
    from app.pptx_export import PPTX_MIME_TYPE, generate_pptx

    title = _bounded_text(arguments.get("title"), field="title", maximum=240)
    _assert_probeable("document.pptx")
    slides = arguments.get("slides")
    try:
        data = generate_pptx(title, slides, content=arguments.get("content"))
    except ValueError as exc:
        raise HTTPException(status_code=422, detail={"code": "pptx_source_invalid"}) from exc
    except Exception as exc:
        _record_failure("document.pptx", "generator_failed")
        raise HTTPException(status_code=502, detail={"code": "pptx_generation_failed"}) from exc
    filename = _safe_filename(arguments.get("filename"), f"{title}.pptx")
    if not filename.lower().endswith(".pptx"):
        filename += ".pptx"
    return _store_artifact(
        data,
        capability_id="document.pptx",
        artifact_type="document.pptx",
        mime_type=PPTX_MIME_TYPE,
        filename=filename,
        title=title,
        principal=principal,
        metadata={"slide_count": len(slides) if isinstance(slides, list) else 1},
    )


async def _generate_project(
    capability_id: str,
    arguments: dict[str, Any],
    principal: ProjectPrincipal,
) -> dict[str, Any]:
    from app.project_builder import (
        ProjectBuildError,
        deterministic_project_zip,
        invoke_project_provider,
        parse_project_payload,
        project_bundle_filename,
    )

    prompt = _bounded_text(arguments.get("prompt"), field="prompt", maximum=20_000)
    try:
        provider_result = await invoke_project_provider(capability_id, prompt)
        payload = parse_project_payload(provider_result.get("content"))
        data = deterministic_project_zip(payload, capability_id)
    except ProjectBuildError as exc:
        _record_failure(capability_id, exc.code)
        raise HTTPException(
            status_code=502,
            detail={
                "code": exc.code,
                "capability": capability_id,
                "retryable": exc.retryable,
            },
        ) from exc

    provider = str(provider_result.get("provider") or "kolibri-provider")[:80]
    model = str(provider_result.get("model") or "kolibri")[:120]
    artifact = _store_artifact(
        data,
        capability_id=capability_id,
        artifact_type="site.bundle" if capability_id == "site.create" else "app.bundle",
        mime_type="application/zip",
        filename=project_bundle_filename(capability_id, payload),
        title=payload.title,
        principal=principal,
        metadata={
            "entrypoint": "index.html",
            "file_count": len(payload.files),
            "provider": provider,
            "model": model,
        },
    )
    # Replace the generic persistence probe with the actual provider identity.
    record_capability_invocation(
        capability_id,
        succeeded=True,
        provider=provider,
        model=model,
        evidence_id=str(artifact["sha256"]),
    )
    preview_url = _signed_preview_url(
        str(artifact["id"]),
        principal.scope_id,
    )
    return {
        "artifact": artifact,
        "preview_url": preview_url,
        "entrypoint": "index.html",
        "files": [
            {
                "path": item.path,
                "size_bytes": len(item.content),
                "sha256": hashlib.sha256(item.content).hexdigest(),
            }
            for item in payload.files
        ],
        "provider": provider,
        "model": model,
    }


async def _generate_or_edit_image(
    capability_id: str,
    arguments: dict[str, Any],
    principal: ProjectPrincipal,
) -> dict[str, Any]:
    """Execute the real image route and return only a verified CAS artifact."""

    from app.image_artifacts import (
        ImageCapabilityUnavailable,
        ImageEditRequest,
        ImageGenerationFailed,
        ImageGenerationRequest,
        edit_image,
        generate_image,
        verify_image_artifact,
    )

    try:
        if capability_id == "image.generate":
            request = ImageGenerationRequest.model_validate(arguments)
            artifact = await generate_image(request, scope_id=principal.scope_id)
        else:
            request = ImageEditRequest.model_validate(arguments)
            artifact = await edit_image(request, scope_id=principal.scope_id)
        artifact = verify_image_artifact(artifact)
        record_capability_invocation(
            capability_id,
            succeeded=True,
            provider="image-provider",
            model=str(artifact.get("model") or "kolibri-image"),
            evidence_id=str(artifact.get("sha256") or "") or None,
        )
    except ValidationError as exc:
        raise HTTPException(
            status_code=422,
            detail={"code": "image_request_invalid", "capability": capability_id},
        ) from exc
    except ImageCapabilityUnavailable as exc:
        _record_failure(capability_id, "image_route_unavailable")
        raise CapabilityInvocationError(
            capability_id,
            "route_not_configured",
            retryable=True,
        ) from exc
    except ImageGenerationFailed as exc:
        _record_failure(capability_id, "image_provider_failed")
        raise HTTPException(
            status_code=502,
            detail={
                "code": "image_execution_failed",
                "capability": capability_id,
                "retryable": True,
            },
        ) from exc
    return {"artifact": artifact}


async def invoke_tool(
    request: ToolInvocationRequest,
    principal: ProjectPrincipal,
) -> dict[str, Any]:
    tool = request.tool.strip().lower()
    if not _TOOL_ID.fullmatch(tool):
        raise HTTPException(status_code=422, detail={"code": "tool_id_invalid"})
    try:
        _assert_probeable(tool)
    except CapabilityInvocationError as exc:
        raise _capability_error(exc) from exc

    if tool == "web.search":
        result = await _web_search(request.arguments)
    elif tool == "document.pdf":
        result = _generate_pdf(request.arguments, principal)
    elif tool == "document.docx":
        result = _generate_docx(request.arguments, principal)
    elif tool == "document.xlsx":
        result = _generate_xlsx(request.arguments, principal)
    elif tool == "document.pptx":
        result = _generate_pptx(request.arguments, principal)
    elif tool in {"image.generate", "image.edit"}:
        result = await _generate_or_edit_image(tool, request.arguments, principal)
    elif tool in {"site.create", "app.create"}:
        result = await _generate_project(tool, request.arguments, principal)
    else:
        # Registered-but-unimplemented capabilities never fall through to an
        # LLM.  The registry composition should already mark them unavailable;
        # this guard also protects against a stale/misconfigured registration.
        _record_failure(tool, "handler_not_implemented")
        raise HTTPException(
            status_code=501,
            detail={"code": "tool_handler_not_implemented", "capability": tool},
        )
    return {"tool": tool, "status": "completed", "result": result}


@router.post("/api/v1/tools/invoke")
@router.post("/v1/tools/invoke", include_in_schema=False)
async def invoke_tool_route(
    request: ToolInvocationRequest,
    principal: ProjectPrincipal = Depends(resolve_project_principal),
):
    return await invoke_tool(request, principal)


def _extract_pdf_text(stored_content: bytes) -> str:
    binary = shutil.which(os.getenv("KOLIBRI_PDFTOTEXT_BINARY", "pdftotext"))
    if not binary:
        _record_failure("file.analyze", "pdf_text_extractor_unavailable")
        raise HTTPException(
            status_code=503,
            detail={
                "code": "pdf_text_extractor_unavailable",
                "message": "Извлечение текста из PDF сейчас недоступно.",
                "retryable": True,
            },
        )
    try:
        with tempfile.NamedTemporaryFile(suffix=".pdf") as source:
            source.write(stored_content)
            source.flush()
            result = subprocess.run(
                [binary, "-enc", "UTF-8", "-nopgbrk", source.name, "-"],
                capture_output=True,
                check=False,
                timeout=20,
            )
    except (OSError, subprocess.TimeoutExpired) as exc:
        _record_failure("file.analyze", "pdf_text_extraction_failed")
        raise ArtifactValidationError("pdf_text_extraction_failed") from exc
    if result.returncode != 0:
        _record_failure("file.analyze", "pdf_text_extraction_failed")
        raise ArtifactValidationError("pdf_text_extraction_failed")
    try:
        text = result.stdout.decode("utf-8")[:_MAX_ANALYSIS_TEXT]
    except UnicodeDecodeError as exc:
        _record_failure("file.analyze", "pdf_text_extraction_failed")
        raise ArtifactValidationError("pdf_text_extraction_failed") from exc
    if not text.strip():
        # An image-only PDF requires an explicit OCR capability.  This handler
        # must never imply that OCR happened when no text layer exists.
        raise ArtifactValidationError("pdf_scanned_or_no_text")
    return text


def _extract_text(stored_content: bytes, mime_type: str) -> str:
    if mime_type.startswith("text/"):
        return stored_content.decode("utf-8")[:_MAX_ANALYSIS_TEXT]
    if mime_type == "application/json":
        parsed = json.loads(stored_content)
        return json.dumps(parsed, ensure_ascii=False, indent=2)[:_MAX_ANALYSIS_TEXT]
    if mime_type == "application/pdf":
        return _extract_pdf_text(stored_content)
    if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        from docx import Document

        document = Document(BytesIO(stored_content))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(parts)[:_MAX_ANALYSIS_TEXT]
    if mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(stored_content), read_only=True, data_only=True)
        try:
            parts: list[str] = []
            for worksheet in workbook.worksheets[:20]:
                parts.append(f"# {worksheet.title}")
                for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                    if index >= 10_000:
                        break
                    parts.append("\t".join("" if value is None else str(value) for value in row))
                    if sum(len(part) for part in parts) >= _MAX_ANALYSIS_TEXT:
                        break
            return "\n".join(parts)[:_MAX_ANALYSIS_TEXT]
        finally:
            workbook.close()
    raise ArtifactValidationError("file_format_not_analyzable")


def _assert_owned(artifact: dict[str, Any], principal: ProjectPrincipal) -> None:
    metadata = artifact.get("metadata") if isinstance(artifact.get("metadata"), dict) else {}
    expected = hashlib.sha256(principal.scope_id.encode()).hexdigest()
    if metadata.get("scope_key") != expected:
        raise ArtifactNotFound("artifact_not_found")


@router.post("/api/v1/files", status_code=201)
@router.post("/v1/files", status_code=201, include_in_schema=False)
async def upload_file(
    file: UploadFile = File(...),
    principal: ProjectPrincipal = Depends(resolve_project_principal),
):
    try:
        _assert_probeable("file.upload")
    except CapabilityInvocationError as exc:
        raise _capability_error(exc) from exc
    data = await file.read(_MAX_UPLOAD_BYTES + 1)
    await file.close()
    if not data or len(data) > _MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail={"code": "file_size_invalid"})
    filename = _safe_filename(file.filename, "upload.bin")
    mime_type = str(file.content_type or "application/octet-stream").lower()
    try:
        artifact = _store_artifact(
            data,
            capability_id="file.upload",
            artifact_type="file.upload",
            mime_type=mime_type,
            filename=filename,
            title=filename,
            principal=principal,
        )
    except HTTPException:
        raise
    return {"object": "file", "artifact": artifact}


@router.post("/api/v1/files/{artifact_id}/analyze", status_code=201)
@router.post("/v1/files/{artifact_id}/analyze", status_code=201, include_in_schema=False)
async def analyze_file(
    artifact_id: str,
    principal: ProjectPrincipal = Depends(resolve_project_principal),
):
    stored = None
    try:
        _assert_probeable("file.analyze")
        stored = get_artifact_store().open(artifact_id)
        _assert_owned(stored.manifest, principal)
        text = _extract_text(stored.content, str(stored.manifest["mime_type"]))
    except CapabilityInvocationError as exc:
        raise _capability_error(exc) from exc
    except ArtifactNotFound as exc:
        raise HTTPException(status_code=404, detail={"code": "file_not_found"}) from exc
    except (ArtifactValidationError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        known_codes = {
            "pdf_scanned_or_no_text",
            "pdf_text_extraction_failed",
            "file_format_not_analyzable",
        }
        code = str(exc) if str(exc) in known_codes else "file_format_not_analyzable"
        raise HTTPException(
            status_code=422,
            detail={
                "code": code,
                "mime_type": stored.manifest.get("mime_type") if stored is not None else None,
                **({"ocr_supported": False} if code == "pdf_scanned_or_no_text" else {}),
            },
        ) from exc
    payload = {
        "source_artifact_id": artifact_id,
        "source_sha256": stored.manifest["sha256"],
        "mime_type": stored.manifest["mime_type"],
        "text": text,
        "characters": len(text),
    }
    artifact = _store_artifact(
        json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8"),
        capability_id="file.analyze",
        artifact_type="file.analysis",
        mime_type="application/json",
        filename=f"analysis-{artifact_id}.json",
        title=f"Анализ: {stored.manifest['title']}",
        principal=principal,
        metadata={"source_artifact_id": artifact_id},
    )
    return {"analysis": payload, "artifact": artifact}


@router.get("/api/v1/files/search")
@router.get("/v1/files/search", include_in_schema=False)
async def search_files(
    q: str = Query(min_length=2, max_length=500),
    limit: int = Query(default=20, ge=1, le=100),
    principal: ProjectPrincipal = Depends(resolve_project_principal),
):
    try:
        _assert_probeable("file.search")
    except CapabilityInvocationError as exc:
        raise _capability_error(exc) from exc
    store = get_artifact_store()
    query = q.casefold().strip()
    matches: list[dict[str, Any]] = []
    # ArtifactStore.list() validates immutable manifests; corrupt entries are
    # not silently presented as search results.
    scope_key = hashlib.sha256(principal.scope_id.encode()).hexdigest()
    for artifact in store.list(metadata={"scope_key": scope_key}):
        if len(matches) >= limit:
            break
        try:
            stored = store.open(str(artifact["id"]))
            text = _extract_text(stored.content, str(artifact["mime_type"]))
        except (ArtifactStoreError, UnicodeDecodeError, json.JSONDecodeError):
            continue
        haystack = f"{artifact.get('title', '')}\n{artifact.get('filename', '')}\n{text}".casefold()
        index = haystack.find(query)
        if index < 0:
            continue
        start = max(0, index - 120)
        end = min(len(haystack), index + len(query) + 220)
        matches.append(
            {
                "artifact": artifact,
                "snippet": haystack[start:end].replace("\n", " "),
            }
        )
    record_capability_invocation(
        "file.search",
        succeeded=True,
        provider="artifact-index",
        evidence_id="sha256:" + json_hash([item["artifact"]["sha256"] for item in matches]),
    )
    return {"query": q, "items": matches, "total": len(matches)}


_PREVIEW_CSP = "; ".join(
    (
        "default-src 'none'",
        "base-uri 'none'",
        "connect-src 'none'",
        "form-action 'none'",
        "object-src 'none'",
        "frame-ancestors 'self'",
        "script-src 'self' 'unsafe-inline'",
        "style-src 'self' 'unsafe-inline'",
        "img-src 'self' data: blob:",
        "font-src 'self' data:",
        "media-src 'self' data: blob:",
        "manifest-src 'self'",
        "worker-src 'none'",
        "sandbox allow-scripts",
    )
)


def _preview_signing_key() -> bytes:
    from app.auth import SECRET_KEY

    secret = (
        os.getenv("KOLIBRI_PREVIEW_SIGNING_SECRET")
        or os.getenv("KOLIBRI_SESSION_SECRET")
        or os.getenv("JWT_SECRET_KEY")
        or SECRET_KEY
    )
    return hashlib.sha256(secret.encode("utf-8")).digest()


def _preview_token_ttl_seconds() -> int:
    try:
        configured = int(os.getenv("KOLIBRI_PREVIEW_TOKEN_TTL_SECONDS", "600"))
    except ValueError:
        configured = 600
    return max(60, min(configured, 3600))


def _urlsafe_encode(value: bytes) -> str:
    return base64.urlsafe_b64encode(value).rstrip(b"=").decode("ascii")


def _urlsafe_decode(value: str) -> bytes:
    return base64.urlsafe_b64decode(value + "=" * (-len(value) % 4))


def _preview_token(
    artifact_id: str,
    scope_key: str,
    *,
    expires_at: int,
) -> str:
    from app.capability_runtime import capability_release_id

    payload = {
        "v": 1,
        "a": artifact_id,
        "s": scope_key,
        "r": capability_release_id(),
        "e": int(expires_at),
    }
    encoded = _urlsafe_encode(
        json.dumps(payload, sort_keys=True, separators=(",", ":")).encode("utf-8")
    )
    signature = _urlsafe_encode(
        hmac.new(_preview_signing_key(), encoded.encode("ascii"), hashlib.sha256).digest()
    )
    return f"{encoded}.{signature}"


def _signed_preview_url(artifact_id: str, scope_id: str) -> str:
    scope_key = hashlib.sha256(scope_id.encode()).hexdigest()
    token = _preview_token(
        artifact_id,
        scope_key,
        expires_at=int(time.time()) + _preview_token_ttl_seconds(),
    )
    return f"/api/v1/previews/{artifact_id}/s/{token}/index.html"


def _assert_preview_token_access(
    token: str,
    artifact_id: str,
    artifact: dict[str, Any],
    *,
    now: int | None = None,
) -> None:
    from app.capability_runtime import capability_release_id

    if not token or len(token) > 2048:
        raise ArtifactNotFound("preview_not_found")
    try:
        encoded, supplied_signature = token.split(".", 1)
        expected_signature = _urlsafe_encode(
            hmac.new(
                _preview_signing_key(),
                encoded.encode("ascii"),
                hashlib.sha256,
            ).digest()
        )
        if not hmac.compare_digest(supplied_signature, expected_signature):
            raise ArtifactNotFound("preview_not_found")
        payload = json.loads(_urlsafe_decode(encoded))
        metadata = artifact.get("metadata")
        scope_key = metadata.get("scope_key") if isinstance(metadata, dict) else None
        current = int(time.time()) if now is None else int(now)
        if (
            not isinstance(payload, dict)
            or payload.get("v") != 1
            or not hmac.compare_digest(str(payload.get("a") or ""), artifact_id)
            or not isinstance(scope_key, str)
            or not hmac.compare_digest(str(payload.get("s") or ""), scope_key)
            or not hmac.compare_digest(
                str(payload.get("r") or ""),
                capability_release_id(),
            )
            or not isinstance(payload.get("e"), int)
            or int(payload["e"]) <= current
        ):
            raise ArtifactNotFound("preview_not_found")
    except (ValueError, TypeError, json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise ArtifactNotFound("preview_not_found") from exc


def _preview_media_type(path: str) -> str:
    suffix = path.rsplit(".", 1)[-1].casefold() if "." in path else ""
    overrides = {
        "css": "text/css",
        "html": "text/html",
        "js": "text/javascript",
        "jsx": "text/javascript",
        "mjs": "text/javascript",
        "svg": "image/svg+xml",
        "ts": "text/plain",
        "tsx": "text/plain",
        "webmanifest": "application/manifest+json",
    }
    return overrides.get(suffix) or mimetypes.guess_type(path)[0] or "text/plain"


def _project_preview_response(
    artifact_id: str,
    asset_path: str,
    principal: ProjectPrincipal | None = None,
    preview_token: str | None = None,
) -> Response:
    from app.project_builder import ProjectBuildError, read_project_archive_file

    try:
        stored = get_artifact_store().open(artifact_id)
        if preview_token is not None:
            _assert_preview_token_access(preview_token, artifact_id, stored.manifest)
        elif principal is not None:
            _assert_owned(stored.manifest, principal)
        else:
            raise ArtifactNotFound("preview_not_found")
        if (
            stored.manifest.get("type") not in {"site.bundle", "app.bundle"}
            or stored.manifest.get("mime_type") != "application/zip"
        ):
            raise ArtifactNotFound("artifact_not_found")
        content = read_project_archive_file(stored.content, asset_path)
    except ArtifactNotFound as exc:
        raise HTTPException(status_code=404, detail={"code": "preview_not_found"}) from exc
    except (ArtifactIntegrityError, ArtifactValidationError) as exc:
        raise HTTPException(status_code=409, detail={"code": "preview_integrity_failed"}) from exc
    except ProjectBuildError as exc:
        status = 404 if exc.code == "preview_asset_not_found" else 422
        raise HTTPException(status_code=status, detail={"code": exc.code}) from exc

    headers = {
        "Content-Security-Policy": _PREVIEW_CSP,
        # The preview document is deliberately sandboxed without
        # allow-same-origin, so its effective origin is opaque.  CORP
        # same-origin would therefore make the browser block its own relative
        # CSS/JS via ORB.  Access remains principal-gated before bytes are
        # returned and CSP still denies network/navigation escape.
        "Cross-Origin-Resource-Policy": "cross-origin",
        "Permissions-Policy": "camera=(), microphone=(), geolocation=(), payment=(), usb=()",
        "Referrer-Policy": "no-referrer",
        "X-Content-Type-Options": "nosniff",
        "X-Frame-Options": "SAMEORIGIN",
        "Cache-Control": "private, no-store",
        "ETag": f'"{hashlib.sha256(content).hexdigest()}"',
    }
    return Response(content=content, media_type=_preview_media_type(asset_path), headers=headers)


@router.get("/api/v1/previews/{artifact_id}")
@router.get("/v1/previews/{artifact_id}", include_in_schema=False)
async def project_preview_index(
    artifact_id: str,
    principal: ProjectPrincipal = Depends(resolve_project_principal),
):
    return _project_preview_response(artifact_id, "index.html", principal)


@router.get("/api/v1/previews/{artifact_id}/s/{preview_token}/{asset_path:path}")
@router.get(
    "/v1/previews/{artifact_id}/s/{preview_token}/{asset_path:path}",
    include_in_schema=False,
)
async def project_preview_signed_asset(
    artifact_id: str,
    preview_token: str,
    asset_path: str,
):
    return _project_preview_response(
        artifact_id,
        asset_path,
        preview_token=preview_token,
    )


@router.get("/api/v1/previews/{artifact_id}/{asset_path:path}")
@router.get("/v1/previews/{artifact_id}/{asset_path:path}", include_in_schema=False)
async def project_preview_asset(
    artifact_id: str,
    asset_path: str,
    principal: ProjectPrincipal = Depends(resolve_project_principal),
):
    return _project_preview_response(artifact_id, asset_path, principal)


def _not_implemented_after_owner_auth(capability_id: str) -> None:
    capability = capability_by_id(capability_id)
    reason = capability.get("reason") if isinstance(capability, dict) else None
    reason_code = reason.get("code") if isinstance(reason, dict) else "capability_not_registered"
    raise _capability_error(CapabilityInvocationError(capability_id, str(reason_code)))


@router.post(
    "/api/v1/integrations/connect",
    dependencies=[Depends(_authorize_api_key_admin)],
)
async def connect_integration(_request: dict[str, Any]):
    """Owner-gated contract; execution stays unavailable until implemented."""

    _not_implemented_after_owner_auth("integration.connect")


@router.post(
    "/api/v1/automations/run",
    dependencies=[Depends(_authorize_api_key_admin)],
)
async def run_automation(_request: dict[str, Any]):
    """Owner-gated contract; never returns a placeholder execution result."""

    _not_implemented_after_owner_auth("automation.run")


__all__ = ["CapabilityInvocationError", "ToolInvocationRequest", "invoke_tool", "router"]
