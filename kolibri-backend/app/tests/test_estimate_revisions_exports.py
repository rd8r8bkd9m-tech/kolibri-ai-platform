import io
import zipfile

import pytest
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.main import app
from app.technology_card import validate_stored_technology_card


@pytest.fixture()
def client():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    testing_session = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        db = testing_session()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        bootstrap = test_client.post("/api/v1/shell/bootstrap")
        assert bootstrap.status_code == 200, bootstrap.text
        yield test_client
    app.dependency_overrides.pop(get_db, None)
    Base.metadata.drop_all(bind=engine)
    engine.dispose()


def estimate_payload(quantity: str = "2.5") -> dict:
    return {
        "title": "Смета на отделочные работы",
        "client": "ООО Заказчик",
        "object_name": "Дом 100 м²",
        "region": "Татарстан",
        "price_as_of": "2026-07-23",
        "currency": "RUB",
        "overhead_rate": "10",
        "vat_rate": "22",
        "sections": [
            {
                "title": "Работы",
                "positions": [
                    {
                        "code": "W-001",
                        "name": "Подготовка стен",
                        "unit": "м²",
                        "quantity": quantity,
                        "price": "100",
                        "source": "Прайс 2026",
                        "comment": "",
                    }
                ],
            }
        ],
    }


def test_clients_objects_crud_and_unnamed_estimate_object(client: TestClient):
    customer = client.post(
        "/api/v1/clients",
        json={
            "name": "ООО Заказчик",
            "inn": "1655000000",
            "kpp": "165501001",
            "address": "Казань, ул. Строителей, 1",
            "contact_person": "Иванов И.И.",
            "phone": "+7 900 000-00-00",
            "email": "client@example.test",
        },
    )
    assert customer.status_code == 201, customer.text
    customer_data = customer.json()
    assert customer_data["inn"] == "1655000000"

    construction_object = client.post(
        "/api/v1/construction-objects",
        json={
            "name": "Дом на Строителей",
            "client_id": customer_data["id"],
            "address": "Казань, ул. Строителей, 1",
            "region": "Татарстан",
        },
    )
    assert construction_object.status_code == 201, construction_object.text
    assert construction_object.json()["client_id"] == customer_data["id"]

    payload = estimate_payload()
    payload.update({"client": "", "object_name": "", "region": ""})
    estimate = client.post("/api/v1/estimates", json=payload)
    assert estimate.status_code == 201, estimate.text
    estimate_data = estimate.json()
    assert estimate_data["object_name"] == "Объект без названия"
    assert estimate_data["object_record_id"]

    objects = client.get("/api/v1/construction-objects")
    assert objects.status_code == 200
    assert any(item["id"] == estimate_data["object_record_id"] for item in objects.json()["items"])


def test_company_profile_requisites_are_saved_once_per_organization(client: TestClient):
    registered = client.post(
        "/api/v1/auth/register",
        json={
            "email": "company-owner@example.test",
            "name": "Владелец",
            "password": "StrongPass-2026",
        },
    )
    assert registered.status_code == 201, registered.text
    organization_list = client.get("/api/v1/organizations")
    assert organization_list.status_code == 200, organization_list.text
    organization = organization_list.json()["items"][0]
    updated = client.patch(
        f"/api/v1/organizations/{organization['id']}/profile",
        json={
            "legal_name": "ООО Подрядчик",
            "inn": "7701000000",
            "director_name": "Петров П.П.",
            "director_title": "Генеральный директор",
            "legal_address": "Москва, ул. Строителей, 1",
        },
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["company_profile"]["legal_name"] == "ООО Подрядчик"


def technology_card() -> dict:
    return validate_stored_technology_card({
        "schema_version": "1.0",
        "title": "Технологическая карта отделочных работ",
        "object_type": "Жилой дом",
        "scope": "Подготовка стен",
        "user_facts": [{"quote": "Подготовка стен", "meaning": "заданный вид работ"}],
        "assumptions": [],
        "exclusions": [],
        "stages": [{
            "sequence": 1,
            "name": "Подготовка",
            "result": "Стены подготовлены",
            "operations": [{
                "sequence": 1,
                "name": "Подготовка стен",
                "method": "Очистить и подготовить основание",
                "prerequisites": [],
                "quality_checks": ["Основание принято"],
                "safety_controls": ["Использовать СИЗ"],
                "resources": [{
                    "kind": "work",
                    "code": "W-001",
                    "name": "Подготовка стен",
                    "unit": "м²",
                    "quantity": "2.5",
                    "price": "100",
                    "quantity_basis": "замер",
                    "procurement_query": "проверить ставку",
                }],
            }],
        }],
    })


def test_estimate_save_recalc_reload_and_immutable_revisions(client: TestClient):
    created_response = client.post("/api/v1/estimates", json=estimate_payload())
    assert created_response.status_code == 201
    assert created_response.headers["etag"] == '"1"'
    created = created_response.json()
    estimate_id = created["id"]
    assert created["version"] == 1
    assert created["price_as_of"] == "2026-07-23"
    assert created["subtotal"] == "250.00"
    assert created["overhead_amount"] == "25.00"
    assert created["vat_amount"] == "60.50"
    assert created["total"] == "335.50"
    assert created["sections"][0]["positions"][0]["sum"] == "250.00"

    revisions = client.get(f"/api/v1/estimates/{estimate_id}/revisions")
    assert revisions.status_code == 200
    assert revisions.json()["total"] == 1
    assert [item["version"] for item in revisions.json()["items"]] == [1]

    missing_precondition = client.put(
        f"/api/v1/estimates/{estimate_id}",
        json={"title": "Blind overwrite"},
    )
    assert missing_precondition.status_code == 428
    assert missing_precondition.json()["detail"]["code"] == "estimate_version_required"

    updated_response = client.put(
        f"/api/v1/estimates/{estimate_id}",
        json={"version": 1, "sections": estimate_payload("3")["sections"]},
    )
    assert updated_response.status_code == 200
    assert updated_response.headers["etag"] == '"2"'
    updated = updated_response.json()
    assert updated["version"] == 2
    assert updated["subtotal"] == "300.00"
    assert updated["overhead_amount"] == "30.00"
    assert updated["vat_amount"] == "72.60"
    assert updated["total"] == "402.60"

    stale = client.put(
        f"/api/v1/estimates/{estimate_id}",
        headers={"If-Match": '"1"'},
        json={"title": "Stale title"},
    )
    assert stale.status_code == 409
    assert stale.headers["etag"] == '"2"'
    assert stale.json()["detail"] == {
        "code": "estimate_version_conflict",
        "message": "Estimate was changed by another save; reload before retrying",
        "expected_version": 1,
        "current_version": 2,
    }

    reloaded = client.get(f"/api/v1/estimates/{estimate_id}")
    assert reloaded.status_code == 200
    assert reloaded.headers["etag"] == '"2"'
    assert reloaded.json() == updated

    missing_recalc_precondition = client.post(
        f"/api/v1/estimates/{estimate_id}/calculate"
    )
    assert missing_recalc_precondition.status_code == 428

    recalculated_response = client.post(
        f"/api/v1/estimates/{estimate_id}/calculate",
        headers={"If-Match": 'W/"2"'},
    )
    assert recalculated_response.status_code == 200
    assert recalculated_response.headers["etag"] == '"3"'
    recalculated = recalculated_response.json()
    assert recalculated["version"] == 3
    assert recalculated["total"] == updated["total"]
    assert recalculated["sections"] == updated["sections"]

    revisions = client.get(f"/api/v1/estimates/{estimate_id}/revisions").json()
    assert revisions["total"] == 3
    assert [item["version"] for item in revisions["items"]] == [3, 2, 1]

    first_revision = client.get(
        f"/api/v1/estimates/{estimate_id}/revisions/1"
    )
    assert first_revision.status_code == 200
    assert first_revision.headers["etag"] == '"1"'
    first_snapshot = first_revision.json()["snapshot"]
    assert first_snapshot["version"] == 1
    assert first_snapshot["total"] == "335.50"
    assert first_snapshot["sections"][0]["positions"][0]["quantity"] == "2.5"

    second_revision = client.get(
        f"/api/v1/estimates/{estimate_id}/revisions/2"
    ).json()["snapshot"]
    assert second_revision["version"] == 2
    assert second_revision["total"] == "402.60"
    assert second_revision["sections"][0]["positions"][0]["quantity"] == "3"


def test_estimate_chat_command_adjustments_and_approved_work_catalog(client: TestClient):
    created = client.post("/api/v1/estimates", json=estimate_payload()).json()
    estimate_id = created["id"]

    changed_response = client.post(
        f"/api/v1/estimates/{estimate_id}/command",
        json={
            "version": 1,
            "command": (
                "Накладные 10%, прибыль 8%, резерв 2%, "
                "генподрядные 3%, скидка 1%, НДС 22%"
            ),
        },
    )
    assert changed_response.status_code == 200, changed_response.text
    changed = changed_response.json()
    assert changed["applied"] is True
    assert changed["estimate"]["version"] == 2
    assert changed["estimate"]["profit_amount"] == "22.00"
    assert changed["estimate"]["contingency_amount"] == "5.94"
    assert changed["estimate"]["general_contractor_amount"] == "9.09"
    assert changed["estimate"]["discount_amount"] == "3.12"
    assert changed["estimate"]["vat_amount"] == "67.96"
    assert changed["estimate"]["total"] == "376.87"

    renamed = client.post(
        f"/api/v1/estimates/{estimate_id}/command",
        json={
            "version": 2,
            "command": "Переименуй Подготовка стен в Подготовка кирпичных стен",
        },
    )
    assert renamed.status_code == 200, renamed.text
    assert renamed.json()["estimate"]["sections"][0]["positions"][0]["name"] == "Подготовка кирпичных стен"

    approved = client.post(
        f"/api/v1/estimates/{estimate_id}/command",
        json={"version": 3, "command": "Утверди смету"},
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()["estimate"]["status"] == "approved"

    catalog = client.get("/api/v1/work-catalog")
    assert catalog.status_code == 200, catalog.text
    assert catalog.json()["total"] == 1
    assert catalog.json()["items"][0]["name"] == "Подготовка кирпичных стен"
    assert catalog.json()["items"][0]["usage_count"] == 1
    item_id = catalog.json()["items"][0]["id"]
    repriced = client.patch(
        f"/api/v1/work-catalog/{item_id}",
        json={"latest_price": "475,50", "price_source": "Ручная цена владельца"},
    )
    assert repriced.status_code == 200, repriced.text
    assert repriced.json()["latest_price"] == "475.50"
    assert repriced.json()["price_source"] == "Ручная цена владельца"
    assert repriced.json()["price_observed_at"] is not None


def test_technology_card_is_hidden_by_default_and_promoted_only_after_approval(client: TestClient):
    payload = estimate_payload()
    payload["technology_card"] = technology_card()
    created_response = client.post("/api/v1/estimates", json=payload)
    assert created_response.status_code == 201, created_response.text
    created = created_response.json()
    estimate_id = created["id"]
    assert "technology_card" not in created

    explicit = client.get(f"/api/v1/estimates/{estimate_id}/technology-card")
    assert explicit.status_code == 200, explicit.text
    assert explicit.json()["technology_card"]["content_sha256"] == payload["technology_card"]["content_sha256"]

    empty_catalog = client.get("/api/v1/technology-catalog")
    assert empty_catalog.status_code == 200
    assert empty_catalog.json()["total"] == 0

    approved = client.post(
        f"/api/v1/estimates/{estimate_id}/command",
        json={"version": 1, "command": "Утверди смету"},
    )
    assert approved.status_code == 200, approved.text
    catalog = client.get("/api/v1/technology-catalog")
    assert catalog.status_code == 200
    assert catalog.json()["total"] == 1
    assert catalog.json()["items"][0]["content_sha256"] == payload["technology_card"]["content_sha256"]


def test_estimate_chat_command_does_not_create_revision_for_unknown_edit(client: TestClient):
    created = client.post("/api/v1/estimates", json=estimate_payload()).json()
    result = client.post(
        f"/api/v1/estimates/{created['id']}/command",
        json={"version": 1, "command": "Сделай красиво"},
    )
    assert result.status_code == 200
    assert result.json()["applied"] is False
    assert result.json()["estimate"]["version"] == 1


def test_revision_backed_pdf_xlsx_and_json_exports(client: TestClient):
    created = client.post("/api/v1/estimates", json=estimate_payload()).json()
    estimate_id = created["id"]
    updated = client.put(
        f"/api/v1/estimates/{estimate_id}",
        headers={"If-Match": '"1"'},
        json={"sections": estimate_payload("3")["sections"]},
    ).json()
    assert updated["version"] == 2

    pdf = client.get(f"/api/v1/estimates/{estimate_id}/pdf?version=1")
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.headers["content-security-policy"] == "default-src 'none'; frame-ancestors 'self'"
    assert pdf.headers["content-disposition"].startswith("inline;")
    assert pdf.headers["cache-control"] == "private, no-store, max-age=0"
    assert pdf.headers["x-content-type-options"] == "nosniff"
    assert pdf.headers["etag"] == '"1"'
    assert pdf.content.startswith(b"%PDF-")
    assert len(pdf.content) > 1_000

    xlsx = client.get(f"/api/v1/estimates/{estimate_id}/export/xlsx?version=2")
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert xlsx.headers["etag"] == '"2"'
    assert xlsx.content.startswith(b"PK\x03\x04")
    assert len(xlsx.content) > 1_000
    with zipfile.ZipFile(io.BytesIO(xlsx.content)) as archive:
        assert archive.testzip() is None
        assert "[Content_Types].xml" in archive.namelist()
        assert "xl/workbook.xml" in archive.namelist()
        assert "xl/worksheets/sheet1.xml" in archive.namelist()
    workbook = load_workbook(io.BytesIO(xlsx.content), read_only=False)
    worksheet = workbook["Смета"]
    assert worksheet["B7"].value == "Не выбран"
    assert worksheet["F12"].data_type == "f"
    assert worksheet["F12"].value == "=ROUND(D12*E12,2)"
    assert worksheet["B12"].alignment.wrap_text is True
    assert worksheet.row_dimensions[12].height >= 18
    assert worksheet.column_dimensions["B"].width >= 48
    assert worksheet.column_dimensions["E"].width >= 30
    assert worksheet.column_dimensions["F"].width >= 19
    assert worksheet.page_setup.fitToHeight == 0
    assert workbook.calculation.calcMode == "auto"
    workbook.close()

    workbook_preview = client.get(
        f"/api/v1/estimates/{estimate_id}/workbook?version=2"
    )
    assert workbook_preview.status_code == 200
    preview_payload = workbook_preview.json()
    assert preview_payload["estimate_id"] == estimate_id
    assert preview_payload["version"] == 2
    assert preview_payload["sheets"][0]["name"] == "Смета"
    assert preview_payload["sheets"][0]["rows"][0]["cells"][0] == created["title"]
    assert any(
        cell == "=ROUND(D12*E12,2)"
        for row in preview_payload["sheets"][0]["rows"]
        for cell in row["cells"]
    )

    first_json = client.get(
        f"/api/v1/estimates/{estimate_id}/export/json?version=1"
    )
    assert first_json.status_code == 200
    assert first_json.headers["content-type"] == "application/json"
    assert first_json.headers["etag"] == '"1"'
    assert first_json.json()["version"] == 1
    assert first_json.json()["total"] == "335.50"
    assert first_json.json()["sections"][0]["positions"][0]["quantity"] == "2.5"

    latest_json = client.get(f"/api/v1/estimates/{estimate_id}/export/json")
    assert latest_json.status_code == 200
    assert latest_json.headers["etag"] == '"2"'
    assert latest_json.json()["version"] == 2
    assert latest_json.json()["total"] == "402.60"


def test_named_tax_regime_controls_vat_and_survives_revision_exports(client: TestClient):
    payload = estimate_payload()
    payload.update({"tax_regime": "usn_vat5", "vat_rate": "22"})
    created = client.post("/api/v1/estimates", json=payload).json()

    assert created["tax_regime"] == "usn_vat5"
    assert created["vat_rate"] == "5"
    assert created["vat_amount"] == "13.75"
    assert created["total"] == "288.75"

    exported = client.get(f"/api/v1/estimates/{created['id']}/export/xlsx?version=1")
    workbook = load_workbook(io.BytesIO(exported.content), read_only=False)
    assert workbook["Смета"]["B7"].value == "УСН · НДС 5%"
    workbook.close()

    updated = client.put(
        f"/api/v1/estimates/{created['id']}",
        headers={"If-Match": '"1"'},
        json={"tax_regime": "npd", "vat_rate": "22"},
    ).json()
    assert updated["tax_regime"] == "npd"
    assert updated["vat_rate"] == "0"
    assert updated["vat_amount"] == "0.00"
