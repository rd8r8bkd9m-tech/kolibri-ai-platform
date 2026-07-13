"""XLSX export — generates Excel files from estimate data."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


def _text(value: object) -> str:
    """Keep user-entered labels from becoming spreadsheet formulas."""
    text = "" if value is None else str(value)
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def _number(value: object, field: str) -> float:
    try:
        number = Decimal(str(value if value is not None else "0"))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"invalid numeric value for {field}") from exc
    if not number.is_finite():
        raise ValueError(f"non-finite numeric value for {field}")
    return float(number)


def generate_estimate_xlsx(estimate: dict) -> bytes:
    """Generate XLSX file from estimate dict."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"

    # Styles
    header_font = Font(name="Arial", bold=True, size=12)
    title_font = Font(name="Arial", bold=True, size=14)
    normal_font = Font(name="Arial", size=11)
    currency = estimate.get("currency", "RUB")
    currency_label = {"RUB": "руб.", "USD": "$", "EUR": "€"}.get(currency, _text(currency))
    money_format = f'#,##0.00 "{currency_label}"'
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )
    header_fill = PatternFill(start_color="E8F8F7", end_color="E8F8F7", fill_type="solid")

    # Title
    ws.merge_cells('A1:F1')
    ws['A1'] = _text(estimate.get('title', 'Смета'))
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center')

    # Info
    row = 3
    info = [
        ('Версия:', estimate.get('version', 1)),
        ('Клиент:', _text(estimate.get('client', ''))),
        ('Объект:', _text(estimate.get('object_name', ''))),
        ('Регион:', _text(estimate.get('region', ''))),
        ('Дата:', _text(estimate.get('created_at', '')[:10] if estimate.get('created_at') else '')),
    ]
    for label, value in info:
        ws[f'A{row}'] = label
        ws[f'A{row}'].font = Font(name="Arial", bold=True, size=11)
        ws[f'B{row}'] = value
        ws[f'B{row}'].font = normal_font
        row += 1

    row += 1

    # Sections
    for section in estimate.get('sections', []):
        ws.merge_cells(f'A{row}:F{row}')
        ws[f'A{row}'] = _text(section.get('title', ''))
        ws[f'A{row}'].font = header_font
        ws[f'A{row}'].fill = header_fill
        row += 1

        # Header
        headers = ['Код', 'Наименование', 'Ед.', 'Кол-во', 'Цена', 'Сумма']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=h)
            cell.font = Font(name="Arial", bold=True, size=10)
            cell.border = thin_border
            cell.fill = header_fill
        row += 1

        # Positions
        for pos in section.get('positions', []):
            ws.cell(row=row, column=1, value=_text(pos.get('code', ''))).font = normal_font
            ws.cell(row=row, column=2, value=_text(pos.get('name', ''))).font = normal_font
            ws.cell(row=row, column=3, value=_text(pos.get('unit', ''))).font = normal_font
            ws.cell(row=row, column=4, value=_number(pos.get('quantity', 0), 'quantity')).font = normal_font
            ws.cell(row=row, column=5, value=_number(pos.get('price', 0), 'price')).font = normal_font
            ws.cell(row=row, column=5).number_format = money_format
            ws.cell(row=row, column=6, value=_number(pos.get('sum', 0), 'sum')).font = normal_font
            ws.cell(row=row, column=6).number_format = money_format
            for col in range(1, 7):
                ws.cell(row=row, column=col).border = thin_border
            row += 1

        # Section subtotal
        ws[f'E{row}'] = 'Итого:'
        ws[f'E{row}'].font = Font(name="Arial", bold=True, size=11)
        ws[f'F{row}'] = _number(section.get('subtotal', 0), 'section subtotal')
        ws[f'F{row}'].font = Font(name="Arial", bold=True, size=11)
        ws[f'F{row}'].number_format = money_format
        row += 2

    # Totals
    ws[f'E{row}'] = 'Подытог:'
    ws[f'E{row}'].font = Font(name="Arial", bold=True, size=11)
    ws[f'F{row}'] = _number(estimate.get('subtotal', 0), 'subtotal')
    ws[f'F{row}'].number_format = money_format
    row += 1

    ws[f'E{row}'] = f'Накладные расходы ({estimate.get("overhead_rate", "0")}%):'
    ws[f'E{row}'].font = Font(name="Arial", bold=True, size=11)
    ws[f'F{row}'] = _number(estimate.get('overhead_amount', 0), 'overhead amount')
    ws[f'F{row}'].number_format = money_format
    row += 1

    ws[f'E{row}'] = f'НДС ({estimate.get("vat_rate", "22")}%):'
    ws[f'E{row}'].font = Font(name="Arial", bold=True, size=11)
    ws[f'F{row}'] = _number(estimate.get('vat_amount', 0), 'vat amount')
    ws[f'F{row}'].number_format = money_format
    row += 1

    ws[f'E{row}'] = 'ИТОГО:'
    ws[f'E{row}'].font = Font(name="Arial", bold=True, size=12)
    ws[f'F{row}'] = _number(estimate.get('total', 0), 'total')
    ws[f'F{row}'].font = Font(name="Arial", bold=True, size=12, color="3ABAB4")
    ws[f'F{row}'].number_format = money_format

    # Column widths
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 8
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 15
    ws.freeze_panes = 'A8'
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = '1:8'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    wb.close()
    content = buf.getvalue()
    if len(content) < 100 or not content.startswith(b"PK\x03\x04"):
        raise RuntimeError("XLSX generator returned an invalid workbook")
    return content
