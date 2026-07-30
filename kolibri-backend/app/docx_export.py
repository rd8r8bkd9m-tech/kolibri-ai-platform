"""DOCX export — generates .docx from document HTML content."""
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO


def html_to_docx(title: str, html: str) -> bytes:
    """Convert HTML content to DOCX bytes."""
    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Arial"
    font.size = Pt(11)

    lines = html.replace("</p>", "\n").replace("<br>", "\n").replace("<br/>", "\n")
    lines = re.sub(r"<h2[^>]*>", "## ", lines)
    lines = re.sub(r"<h3[^>]*>", "### ", lines)
    lines = re.sub(r"<strong>", "**", lines)
    lines = re.sub(r"</strong>", "**", lines)
    lines = re.sub(r"<[^>]+>", "", lines)
    lines = lines.replace("&nbsp;", " ").replace("&mdash;", "—").replace("&laquo;", "«").replace("&raquo;", "»")

    for line in lines.split("\n"):
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
            continue

        if stripped.startswith("## "):
            heading = doc.add_heading(stripped[3:].strip(), level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
        else:
            p = doc.add_paragraph()
            parts = stripped.split("**")
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 == 1:
                    run.bold = True

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
