from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _font_name() -> str:
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            name = "KolibriSans"
            try:
                pdfmetrics.registerFont(TTFont(name, str(path)))
            except Exception:
                pass
            return name
    return "Helvetica"


def render_pdf(estimate: dict[str, Any]) -> bytes:
    buffer = io.BytesIO()
    font = _font_name()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=18*mm, bottomMargin=18*mm)
    styles = getSampleStyleSheet()
    for style_name in ["Title", "Heading2", "BodyText"]:
        styles[style_name].fontName = font
    story = [
        Paragraph("Коммерческое предложение", styles["Title"]),
        Spacer(1, 5*mm),
        Paragraph(estimate.get("title", "Смета"), styles["Heading2"]),
        Paragraph(f"Клиент: {estimate.get('client_name') or 'не указан'}", styles["BodyText"]),
        Paragraph(f"Регион: {estimate.get('region') or 'не указан'}", styles["BodyText"]),
        Paragraph(f"Статус: {estimate.get('status')}", styles["BodyText"]),
        Spacer(1, 5*mm),
    ]
    rows = [["Раздел", "Работа", "Ед.", "Кол-во", "Цена", "Сумма"]]
    for item in estimate.get("items", []):
        rows.append([
            item.get("section", ""), item.get("name", ""), item.get("unit", ""),
            item.get("quantity", "0"), item.get("unit_price", "0"), item.get("line_total", "0"),
        ])
    rows.append(["", "", "", "", "Итого", estimate.get("total", "0")])
    table = Table(rows, colWidths=[28*mm, 55*mm, 15*mm, 18*mm, 23*mm, 25*mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0,0), (-1,-1), font),
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#182522")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#D6D0C5")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("ALIGN", (3,1), (-1,-1), "RIGHT"),
        ("BACKGROUND", (0,-1), (-1,-1), colors.HexColor("#F2EBDD")),
        ("FONTNAME", (4,-1), (-1,-1), font),
    ]))
    story += [table, Spacer(1, 6*mm)]
    if estimate.get("sources"):
        story.append(Paragraph("Источники цен", styles["Heading2"]))
        for source in estimate["sources"]:
            story.append(Paragraph(
                f"{source['title']} · {source['region']} · {source['price_date']} · {source['verification_status']} · {source['url']}",
                styles["BodyText"],
            ))
    doc.build(story)
    return buffer.getvalue()


def render_xlsx(estimate: dict[str, Any]) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Смета"
    headers = ["Раздел", "Работа", "Ед.", "Количество", "Цена", "Коэффициент", "Сумма", "Источник"]
    ws.append([estimate.get("title")]); ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers))
    ws["A1"].font = Font(bold=True, size=16); ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(headers)
    for cell in ws[2]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="182522")
    for item in estimate.get("items", []):
        source=item.get("source") or {}
        ws.append([item.get("section"),item.get("name"),item.get("unit"),float(item.get("quantity",0)),float(item.get("unit_price",0)),float(item.get("coefficient",1)),float(item.get("line_total",0)),source.get("url","")])
    ws.append(["", "", "", "", "", "Итого", float(estimate.get("total",0)), ""])
    for col, width in {"A":20,"B":38,"C":12,"D":14,"E":16,"F":14,"G":16,"H":45}.items(): ws.column_dimensions[col].width=width
    buffer=io.BytesIO(); wb.save(buffer); return buffer.getvalue()


def render_docx(estimate: dict[str, Any]) -> bytes:
    doc=Document(); doc.add_heading("Коммерческое предложение",0); doc.add_heading(estimate.get("title","Смета"),1)
    doc.add_paragraph(f"Клиент: {estimate.get('client_name') or 'не указан'}")
    doc.add_paragraph(f"Регион: {estimate.get('region') or 'не указан'}")
    table=doc.add_table(rows=1, cols=6); table.style="Table Grid"
    for cell, title in zip(table.rows[0].cells,["Раздел","Работа","Ед.","Кол-во","Цена","Сумма"]): cell.text=title
    for item in estimate.get("items",[]):
        cells=table.add_row().cells
        values=[item.get("section",""),item.get("name",""),item.get("unit",""),item.get("quantity","0"),item.get("unit_price","0"),item.get("line_total","0")]
        for cell,value in zip(cells,values): cell.text=str(value)
    doc.add_heading(f"Итого: {estimate.get('total','0')} ₽",1)
    if estimate.get("status")!="verified": doc.add_paragraph("Предварительный расчёт. Требуется подтверждение исходных данных и цен.")
    buffer=io.BytesIO(); doc.save(buffer); return buffer.getvalue()


def render_json(estimate: dict[str, Any]) -> bytes:
    return json.dumps(estimate, ensure_ascii=False, indent=2).encode("utf-8")


def render_markdown(estimate: dict[str, Any]) -> bytes:
    lines=[f"# {estimate.get('title','Смета')}","",f"Статус: **{estimate.get('status')}**",f"Итого: **{estimate.get('total')} ₽**","","## Допущения"]
    if estimate.get("status")!="verified": lines.append("- Итог является предварительным до подтверждения источников и исходных данных.")
    else: lines.append("- Все строки связаны с проверенными источниками цен.")
    lines += ["","## Источники"]
    for source in estimate.get("sources",[]): lines.append(f"- [{source['title']}]({source['url']}) — {source['region']}, {source['price_date']}, {source['verification_status']}")
    return "\n".join(lines).encode("utf-8")
