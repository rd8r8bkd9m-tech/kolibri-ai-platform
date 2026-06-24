import os
import json
import re
import sqlite3
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Optional, Dict, Any
from urllib.parse import quote
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable,
    KeepTogether,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from config import DATA_DIR, DB_PATH
from estimates import normalize_estimate


DOCS_DIR = DATA_DIR / "documents"
DOCS_DIR.mkdir(parents=True, exist_ok=True)

COUNTER_PATH = DOCS_DIR / "_counters.json"
DOCUMENT_INDEX_PATH = DOCS_DIR / "_documents.json"


def _safe_text(value: Any, default: str = "") -> str:
    if value is None:
        return default
    return escape(_strip_unsupported_pdf_symbols(str(value)))


def _strip_unsupported_pdf_symbols(value: str) -> str:
    cleaned = []
    for char in value:
        if char in {"\ufe0e", "\ufe0f", "⚠", "✅", "❗", "❌", "☑", "☐", "□"}:
            continue
        category = unicodedata.category(char)
        if category in {"Cc", "Cf", "Co", "Cs"}:
            if char in {"\n", "\t"}:
                cleaned.append(" ")
            continue
        if category == "So":
            continue
        cleaned.append(char)
    return re.sub(r"\s+", " ", "".join(cleaned)).strip()


def _money(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _document_meta(payload: dict | None, prefix: str, stem: str, ext: str = "pdf") -> tuple[str, datetime, str]:
    meta = payload if isinstance(payload, dict) else {}
    doc_number = meta.get("number") or meta.get("doc_number") or _next_doc_number(prefix)
    raw_date = meta.get("date")
    if raw_date:
        try:
            doc_date = datetime.fromisoformat(str(raw_date).replace("Z", "+00:00"))
        except ValueError:
            doc_date = datetime.now()
    else:
        doc_date = datetime.now()
    filename = meta.get("filename") or f"{stem}_{str(doc_number).replace('/', '-')}.{ext}"
    return str(doc_number), doc_date, filename


def _load_document_index() -> dict:
    try:
        if DOCUMENT_INDEX_PATH.exists():
            data = json.loads(DOCUMENT_INDEX_PATH.read_text(encoding="utf-8"))
            if isinstance(data, dict):
                return data
    except Exception:
        pass
    return {}


def _save_document_index(index: dict) -> None:
    try:
        DOCUMENT_INDEX_PATH.write_text(json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass


def _extract_estimate_title(payload: Any) -> str:
    if not isinstance(payload, dict):
        return ""
    root = payload.get("смета") if isinstance(payload.get("смета"), dict) else payload
    estimate_meta = root.get("estimate") if isinstance(root.get("estimate"), dict) else {}
    for value in (
        estimate_meta.get("title"),
        root.get("title"),
        root.get("проект"),
        root.get("name"),
        root.get("название"),
    ):
        if value:
            return str(value).strip()
    obj = root.get("object") if isinstance(root.get("object"), dict) else root.get("объект") if isinstance(root.get("объект"), dict) else {}
    if isinstance(obj, dict) and obj.get("name"):
        return f"Смета: {obj['name']}"
    return ""


def _extract_json_payload(text: str) -> dict | None:
    value = str(text or "").strip()
    if not value:
        return None
    candidates = [value]
    start = value.find("{")
    end = value.rfind("}")
    if start != -1 and end > start:
        candidates.append(value[start:end + 1])
    for candidate in candidates:
        try:
            parsed = json.loads(candidate)
            if isinstance(parsed, dict):
                return parsed
        except Exception:
            continue
    return None


def _estimate_title_from_messages(created_ts: float) -> str:
    payload = _estimate_payload_from_messages(created_ts)
    return _extract_estimate_title(payload) if payload else ""


def _estimate_payload_from_messages(created_ts: float) -> dict | None:
    try:
        if not DB_PATH.exists():
            return None
        with sqlite3.connect(DB_PATH) as conn:
            rows = conn.execute(
                """
                SELECT content, created_at
                FROM messages
                WHERE role = 'assistant'
                  AND content LIKE '%{%'
                  AND created_at <= ?
                ORDER BY created_at DESC
                LIMIT 40
                """,
                (created_ts + 5,),
            ).fetchall()
    except Exception:
        return None

    for content, _created_at in rows:
        payload = _extract_json_payload(content)
        if payload:
            return payload
    return None


def _fallback_document_title(filename: str) -> str:
    stem = Path(filename).stem
    if stem.startswith("estimate_"):
        number = stem.replace("estimate_", "").replace("-", "/")
        return f"Смета {number}".strip()
    if stem.startswith("kp_"):
        return "Коммерческое предложение"
    if stem.startswith("act_"):
        return "Акт выполненных работ"
    if stem.startswith("invoice_"):
        return "Счёт на оплату"
    return re.sub(r"[_-]+", " ", stem).strip().capitalize() or filename


def _document_number_from_filename(filename: str) -> str:
    stem = Path(filename).stem
    if "_" not in stem:
        return ""
    return stem.split("_", 1)[1].replace("-", "/")


def _document_type(filename: str) -> str:
    stem = Path(filename).stem.lower()
    if stem.startswith("estimate_"):
        return "estimate"
    if stem.startswith("kp_"):
        return "commercial_offer"
    if stem.startswith("act_"):
        return "act"
    if stem.startswith("invoice_"):
        return "invoice"
    return "document"


def _record_document(path: str, *, kind: str, title: str, doc_number: str, doc_date: datetime, estimate: dict | None = None) -> None:
    file_path = Path(path)
    index = _load_document_index()
    totals = estimate.get("totals", {}) if isinstance(estimate, dict) else {}
    index[file_path.name] = {
        "filename": file_path.name,
        "title": title or _fallback_document_title(file_path.name),
        "type": kind,
        "estimate": estimate if isinstance(estimate, dict) and kind == "estimate" else None,
        "document_number": doc_number,
        "document_date": doc_date.isoformat(),
        "created_at": datetime.fromtimestamp(file_path.stat().st_ctime).isoformat() if file_path.exists() else datetime.now().isoformat(),
        "total": totals.get("grand_total") if isinstance(totals, dict) else None,
        "items_count": len(estimate.get("items", [])) if isinstance(estimate, dict) else None,
        "sections_count": len(estimate.get("sections", [])) if isinstance(estimate, dict) else None,
}
    _save_document_index(index)


def get_document_estimate(filename: str) -> dict | None:
    if "/" in filename or ".." in filename:
        return None
    file_path = DOCS_DIR / filename
    if not file_path.exists():
        return None

    index = _load_document_index()
    meta = index.get(file_path.name, {}) if isinstance(index.get(file_path.name), dict) else {}
    raw_estimate = meta.get("estimate") if isinstance(meta.get("estimate"), dict) else None
    if not raw_estimate:
        raw_estimate = _estimate_payload_from_messages(file_path.stat().st_ctime)

    estimate = normalize_estimate(raw_estimate or {})
    fallback_title = meta.get("title") or _estimate_title_from_messages(file_path.stat().st_ctime) or _fallback_document_title(file_path.name)
    if fallback_title and estimate.get("title") in ("", "Смета"):
        estimate["title"] = fallback_title
        estimate.setdefault("estimate", {})["title"] = fallback_title
    estimate["source_document"] = {
        "filename": file_path.name,
        "url": f"/api/documents/file/{quote(file_path.name)}",
        "document_number": meta.get("document_number") or _document_number_from_filename(file_path.name),
        "created_at": meta.get("created_at") or datetime.fromtimestamp(file_path.stat().st_ctime).isoformat(),
    }
    return estimate


# ---------------------------------------------------------------------------
# Fonts
# ---------------------------------------------------------------------------

def _register_fonts():
    font_paths = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/ubuntu/Ubuntu-R.ttf",
    ]
    for fp in font_paths:
        if os.path.exists(fp):
            try:
                pdfmetrics.registerFont(TTFont("KolibriFont", fp))
                return "KolibriFont"
            except Exception:
                continue
    return "Helvetica"


def _register_bold_fonts():
    font_paths = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
        "/usr/share/fonts/truetype/ubuntu/Ubuntu-B.ttf",
    ]
    for fp in font_paths:
        if os.path.exists(fp):
            try:
                pdfmetrics.registerFont(TTFont("KolibriFontBold", fp))
                return "KolibriFontBold"
            except Exception:
                continue
    return "Helvetica-Bold"


FONT_NAME = _register_fonts()
FONT_BOLD = _register_bold_fonts()

# Brand colors
COL_PRIMARY = colors.HexColor("#1a202c")
COL_ACCENT = colors.HexColor("#2b6cb0")
COL_HEADER_BG = colors.HexColor("#2d3748")
COL_ROW_ALT = colors.HexColor("#f7fafc")
COL_BORDER = colors.HexColor("#cbd5e0")
COL_GREY = colors.HexColor("#718096")
COL_LIGHT_GREY = colors.HexColor("#e2e8f0")


# ---------------------------------------------------------------------------
# Document numbering
# ---------------------------------------------------------------------------

def _next_doc_number(prefix: str) -> str:
    counters: Dict[str, int] = {}
    if COUNTER_PATH.exists():
        try:
            counters = json.loads(COUNTER_PATH.read_text())
        except Exception:
            pass
    current = counters.get(prefix, 0) + 1
    counters[prefix] = current
    COUNTER_PATH.write_text(json.dumps(counters, ensure_ascii=False, indent=2))
    return f"{prefix}-{current:03d}"


# ---------------------------------------------------------------------------
# City extraction
# ---------------------------------------------------------------------------

RU_MONTHS = {
    1: "января", 2: "февраля", 3: "марта", 4: "апреля", 5: "мая", 6: "июня",
    7: "июля", 8: "августа", 9: "сентября", 10: "октября", 11: "ноября", 12: "декабря",
}


def _ru_date_long(dt: datetime) -> str:
    return f"«{dt.day}» {RU_MONTHS[dt.month]} {dt.year} г."


def _extract_city(address: str) -> str:
    if not address:
        return ""
    for token in address.split(","):
        token = token.strip()
        if token.startswith("г."):
            return token[2:].strip()
        if token.startswith("город"):
            return token.replace("город", "").strip()
    return ""


# ---------------------------------------------------------------------------
# Styles factory
# ---------------------------------------------------------------------------

def _make_styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("KTitle", parent=base["Title"], fontName=FONT_BOLD, fontSize=18, spaceAfter=4, textColor=COL_PRIMARY),
        "subtitle": ParagraphStyle("KSub", parent=base["Heading2"], fontName=FONT_BOLD, fontSize=13, spaceAfter=6, textColor=COL_PRIMARY),
        "normal": ParagraphStyle("KNorm", parent=base["Normal"], fontName=FONT_NAME, fontSize=10, leading=14),
        "small": ParagraphStyle("KSmall", parent=base["Normal"], fontName=FONT_NAME, fontSize=8, textColor=COL_GREY, leading=10),
        "bold": ParagraphStyle("KBold", parent=base["Normal"], fontName=FONT_BOLD, fontSize=10, leading=14),
        "right": ParagraphStyle("KRight", parent=base["Normal"], fontName=FONT_NAME, fontSize=10, alignment=2),
        "header_cell": ParagraphStyle("KHCell", parent=base["Normal"], fontName=FONT_BOLD, fontSize=9, textColor=colors.white, leading=12),
        "cell": ParagraphStyle("KCell", parent=base["Normal"], fontName=FONT_NAME, fontSize=9, leading=12),
        "cell_bold": ParagraphStyle("KCellB", parent=base["Normal"], fontName=FONT_BOLD, fontSize=9, leading=12),
        "section_row": ParagraphStyle("KSect", parent=base["Normal"], fontName=FONT_BOLD, fontSize=9, textColor=COL_ACCENT, leading=12),
        "total_label": ParagraphStyle("KTotL", parent=base["Normal"], fontName=FONT_BOLD, fontSize=10, alignment=2, leading=14),
        "total_value": ParagraphStyle("KTotV", parent=base["Normal"], fontName=FONT_BOLD, fontSize=10, leading=14),
        "grand_label": ParagraphStyle("KGrandL", parent=base["Normal"], fontName=FONT_BOLD, fontSize=12, alignment=2, leading=16),
        "grand_value": ParagraphStyle("KGrandV", parent=base["Normal"], fontName=FONT_BOLD, fontSize=12, textColor=COL_ACCENT, leading=16),
    }


# ---------------------------------------------------------------------------
# Page template with header + footer + page numbers
# ---------------------------------------------------------------------------

def _make_page_template(doc_title: str, doc_number: str, company: dict = None, doc_date: datetime | None = None):
    def _on_page(canvas, doc):
        canvas.saveState()
        w, h = A4

        # Header line
        canvas.setStrokeColor(COL_ACCENT)
        canvas.setLineWidth(1.5)
        canvas.line(20 * mm, h - 15 * mm, w - 20 * mm, h - 15 * mm)

        # Header text
        canvas.setFont(FONT_NAME, 7)
        canvas.setFillColor(COL_GREY)
        comp_name = (company or {}).get("name", "")
        header_left = comp_name if comp_name else doc_title
        canvas.drawString(20 * mm, h - 13 * mm, header_left)
        canvas.drawRightString(w - 20 * mm, h - 13 * mm, doc_number)

        # Footer line
        canvas.setStrokeColor(COL_LIGHT_GREY)
        canvas.setLineWidth(0.5)
        canvas.line(20 * mm, 15 * mm, w - 20 * mm, 15 * mm)

        # Footer text
        canvas.setFont(FONT_NAME, 7)
        canvas.setFillColor(COL_GREY)
        canvas.drawString(20 * mm, 10 * mm, (doc_date or datetime.now()).strftime("%d.%m.%Y"))
        canvas.drawRightString(w - 20 * mm, 10 * mm, f"Стр. {doc.page}")

        # Company footer details
        if company:
            canvas.setFont(FONT_NAME, 6)
            footer_parts = []
            if company.get("inn"):
                footer_parts.append(f"ИНН {company['inn']}")
            if company.get("phone"):
                footer_parts.append(company["phone"])
            if company.get("email"):
                footer_parts.append(company["email"])
            if footer_parts:
                canvas.drawCentredString(w / 2, 6 * mm, "  |  ".join(footer_parts))

        canvas.restoreState()

    return _on_page


# ---------------------------------------------------------------------------
# Shared table builder
# ---------------------------------------------------------------------------

def _build_items_table(items: list, totals: dict, st: dict, show_sections: bool = False, sections: list | None = None):
    col_widths = [25, 165, 38, 45, 82, 90]

    # Header row
    header = [
        Paragraph("<b>№</b>", st["header_cell"]),
        Paragraph("<b>Наименование</b>", st["header_cell"]),
        Paragraph("<b>Ед.</b>", st["header_cell"]),
        Paragraph("<b>Кол-во</b>", st["header_cell"]),
        Paragraph("<b>Цена</b>", st["header_cell"]),
        Paragraph("<b>Сумма</b>", st["header_cell"]),
    ]
    data = [header]
    row_types = ["header"]  # track row types for styling

    if show_sections:
        groups = []
        seen_item_ids = set()
        for section in sorted(sections or [], key=lambda item: item.get("number", 0)):
            section_id = section.get("id")
            section_items = [item for item in items if item.get("section_id") == section_id]
            if not section_items:
                continue
            seen_item_ids.update(id(item) for item in section_items)
            groups.append((_safe_text(section.get("name", "Раздел")), section_items))
        orphan_items = [item for item in items if id(item) not in seen_item_ids]
        if orphan_items:
            by_type = [
                ("Работы", [i for i in orphan_items if i.get("type") == "work"]),
                ("Материалы", [i for i in orphan_items if i.get("type") == "material"]),
                ("Услуги", [i for i in orphan_items if i.get("type") == "service"]),
            ]
            groups.extend((name, group_items) for name, group_items in by_type if group_items)
        if not groups:
            groups = [("", items)]
    else:
        groups = [("", items)]

    idx = 0
    for section_name, section_items in groups:
        if section_name:
            data.append([
                Paragraph(f"<b>{section_name}</b>", st["section_row"]),
                "", "", "", "", "",
            ])
            row_types.append("section")

        for item in section_items:
            idx += 1
            data.append([
                Paragraph(str(idx), st["cell"]),
                Paragraph(_safe_text(item.get("name", "")), st["cell"]),
                Paragraph(_safe_text(item.get("unit", "")), st["cell"]),
                Paragraph(_safe_text(item.get("quantity", 0)), st["cell"]),
                Paragraph(f"{_money(item.get('unit_price')):,.2f}", st["cell"]),
                Paragraph(f"{_money(item.get('total')):,.2f}", st["cell_bold"]),
            ])
            row_types.append("item")

    # Totals rows
    works_total = _money(totals.get("works", totals.get("works_total", 0)))
    mats_total = _money(totals.get("materials", totals.get("materials_total", 0)))
    services_total = _money(totals.get("services", totals.get("services_total", 0)))
    delivery = _money(totals.get("delivery", 0))
    discount = _money(totals.get("discount", 0))
    grand = _money(totals.get("grand_total", 0))

    totals_data = []
    totals_row_types = []
    if show_sections:
        if works_total > 0:
            totals_data.append(["", "", "", "", Paragraph("Работы:", st["total_label"]), Paragraph(f"{works_total:,.2f}", st["total_value"])])
            totals_row_types.append("subtotal")
        if mats_total > 0:
            totals_data.append(["", "", "", "", Paragraph("Материалы:", st["total_label"]), Paragraph(f"{mats_total:,.2f}", st["total_value"])])
            totals_row_types.append("subtotal")
        if services_total > 0:
            totals_data.append(["", "", "", "", Paragraph("Услуги:", st["total_label"]), Paragraph(f"{services_total:,.2f}", st["total_value"])])
            totals_row_types.append("subtotal")
    if delivery > 0:
        totals_data.append(["", "", "", "", Paragraph("Доставка:", st["total_label"]), Paragraph(f"{delivery:,.2f}", st["total_value"])])
        totals_row_types.append("subtotal")
    if discount > 0:
        totals_data.append(["", "", "", "", Paragraph("Скидка:", st["total_label"]), Paragraph(f"-{discount:,.2f}", st["total_value"])])
        totals_row_types.append("subtotal")

    # Grand total
    totals_data.append(["", "", "", "", Paragraph("ИТОГО:", st["grand_label"]), Paragraph(f"{grand:,.2f} руб.", st["grand_value"])])
    totals_row_types.append("grand")

    table = Table(data, colWidths=col_widths, repeatRows=1)

    style_cmds = [
        ("FONTNAME", (0, 0), (-1, -1), FONT_NAME),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        # Header
        ("BACKGROUND", (0, 0), (-1, 0), COL_HEADER_BG),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
        # Grid for items
        ("GRID", (0, 0), (-1, -1), 0.5, COL_BORDER),
    ]

    # Alternating row colors for item rows
    item_row_idx = 0
    for i, rtype in enumerate(row_types):
        if rtype == "item":
            if item_row_idx % 2 == 1:
                style_cmds.append(("BACKGROUND", (0, i), (-1, i), COL_ROW_ALT))
            item_row_idx += 1
        elif rtype == "section":
            style_cmds.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#edf2f7")))
            style_cmds.append(("SPAN", (0, i), (-1, i)))

    table.setStyle(TableStyle(style_cmds))

    totals_table = Table(totals_data, colWidths=col_widths, repeatRows=0)
    totals_style_cmds = [
        ("FONTNAME", (0, 0), (-1, -1), FONT_NAME),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("ALIGN", (3, 0), (-1, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.5, COL_BORDER),
    ]
    for i, rtype in enumerate(totals_row_types):
        if rtype == "grand":
            totals_style_cmds.append(("LINEABOVE", (0, i), (-1, i), 2, COL_ACCENT))
            totals_style_cmds.append(("BACKGROUND", (4, i), (-1, i), colors.HexColor("#ebf8ff")))
            totals_style_cmds.append(("TOPPADDING", (0, i), (-1, i), 8))
            totals_style_cmds.append(("BOTTOMPADDING", (0, i), (-1, i), 8))
    totals_table.setStyle(TableStyle(totals_style_cmds))

    return [table, KeepTogether([totals_table])]


# ---------------------------------------------------------------------------
# Company header block
# ---------------------------------------------------------------------------

def _company_header_block(company: dict, st: dict) -> list:
    elements = []
    name = company.get("name", "")
    if name:
        elements.append(Paragraph(f"<b>{name}</b>", st["subtitle"]))
    details = []
    if company.get("inn"):
        details.append(f"ИНН: {company['inn']}")
    if company.get("kpp"):
        details.append(f"КПП: {company['kpp']}")
    if company.get("phone"):
        details.append(f"Тел: {company['phone']}")
    if company.get("email"):
        details.append(f"Email: {company['email']}")
    if details:
        elements.append(Paragraph(" &nbsp;&nbsp;|&nbsp;&nbsp; ".join(details), st["small"]))
    if company.get("address"):
        elements.append(Paragraph(company["address"], st["small"]))
    return elements


# ---------------------------------------------------------------------------
# Client info block
# ---------------------------------------------------------------------------

def _client_block(client: dict, st: dict) -> list:
    elements = []
    if client.get("name"):
        elements.append(Paragraph(f"<b>Заказчик:</b> {_safe_text(client['name'])}", st["normal"]))
    if client.get("phone"):
        elements.append(Paragraph(f"<b>Телефон:</b> {_safe_text(client['phone'])}", st["normal"]))
    if client.get("address"):
        elements.append(Paragraph(f"<b>Адрес:</b> {_safe_text(client['address'])}", st["normal"]))
    return elements


def _canonical_estimate(payload: dict) -> dict:
    if isinstance(payload, dict) and isinstance(payload.get("estimate"), dict) and not any(
        key in payload for key in ("items", "sections", "client", "object", "смета")
    ):
        payload = payload["estimate"]
    return normalize_estimate(payload if isinstance(payload, dict) else {})


def _object_block(obj: dict, estimate: dict, st: dict) -> list:
    elements = []
    parts = []
    if obj.get("name"):
        parts.append(_safe_text(obj["name"]))
    if obj.get("type") and obj.get("type") != obj.get("name"):
        parts.append(_safe_text(obj["type"]))
    if obj.get("area"):
        parts.append(f"{_safe_text(obj['area'])} м²")
    if parts:
        elements.append(Paragraph(f"<b>Объект:</b> {', '.join(parts)}", st["normal"]))
    if obj.get("address"):
        elements.append(Paragraph(f"<b>Адрес объекта:</b> {_safe_text(obj['address'])}", st["normal"]))
    if estimate.get("region"):
        elements.append(Paragraph(f"<b>Регион:</b> {_safe_text(estimate['region'])}", st["normal"]))
    return elements


# ---------------------------------------------------------------------------
# Estimate PDF
# ---------------------------------------------------------------------------

def generate_estimate_pdf(estimate: dict, company: dict = None, output_path: Optional[str] = None, document: dict = None) -> str:
    estimate = _canonical_estimate(estimate)
    doc_number, doc_date, filename = _document_meta(document, "СМТ", "estimate")
    if not output_path:
        output_path = str(DOCS_DIR / filename)

    st = _make_styles()
    on_page = _make_page_template("СМЕТА", doc_number, company, doc_date)

    doc = SimpleDocTemplate(
        output_path, pagesize=A4,
        topMargin=22 * mm, bottomMargin=22 * mm,
        leftMargin=20 * mm, rightMargin=20 * mm,
    )

    elements = []

    # Title
    elements.append(Paragraph(_safe_text(estimate.get("title", "Смета")), st["title"]))
    elements.append(Paragraph(f"№ {_safe_text(doc_number)} от {doc_date.strftime('%d.%m.%Y')}", st["small"]))
    elements.append(Spacer(1, 4 * mm))
    elements.append(HRFlowable(width="100%", thickness=1, color=COL_LIGHT_GREY))
    elements.append(Spacer(1, 4 * mm))

    # Company header
    if company:
        elements.extend(_company_header_block(company, st))
        elements.append(Spacer(1, 3 * mm))

    # Client info
    client = estimate.get("client", {})
    elements.extend(_client_block(client, st))

    # Object info
    obj = estimate.get("object", {})
    object_elements = _object_block(obj, estimate, st)
    if object_elements:
        elements.append(Spacer(1, 2 * mm))
        elements.extend(object_elements)

    elements.append(Spacer(1, 6 * mm))

    # Items table
    items = estimate.get("items", [])
    totals = estimate.get("totals", {})
    if items:
        elements.extend(_build_items_table(items, totals, st, show_sections=True, sections=estimate.get("sections", [])))
        elements.append(Spacer(1, 6 * mm))

    # Assumptions
    assumptions = estimate.get("assumptions", [])
    if assumptions:
        elements.append(Paragraph("<b>Допущения:</b>", st["bold"]))
        for a in assumptions:
            elements.append(Paragraph(f"&bull; {_safe_text(a)}", st["small"]))
        elements.append(Spacer(1, 3 * mm))

    questions = estimate.get("questions", [])
    if questions:
        elements.append(Paragraph("<b>Вопросы:</b>", st["bold"]))
        for q in questions:
            elements.append(Paragraph(f"&bull; {_safe_text(q)}", st["small"]))
        elements.append(Spacer(1, 3 * mm))

    # Warnings
    warnings = estimate.get("warnings", [])
    if warnings:
        elements.append(Paragraph("<b>Примечания:</b>", st["bold"]))
        for w in warnings:
            elements.append(Paragraph(f"&bull; {_safe_text(w)}", st["small"]))

    # Footer
    elements.append(Spacer(1, 8 * mm))
    elements.append(HRFlowable(width="100%", thickness=0.5, color=COL_LIGHT_GREY))
    elements.append(Spacer(1, 3 * mm))
    versions = estimate.get("versions") or []
    version = versions[-1].get("number") if versions and isinstance(versions[-1], dict) else estimate.get("version", 1)
    elements.append(Paragraph(f"Версия: v{_safe_text(version)}", st["small"]))

    doc.build(elements, onFirstPage=on_page, onLaterPages=on_page)
    _record_document(
        output_path,
        kind="estimate",
        title=estimate.get("title", "Смета"),
        doc_number=doc_number,
        doc_date=doc_date,
        estimate=estimate,
    )
    return output_path


# ---------------------------------------------------------------------------
# Commercial Offer (КП) PDF
# ---------------------------------------------------------------------------

def generate_commercial_offer_pdf(estimate: dict, company: dict = None, output_path: Optional[str] = None) -> str:
    estimate = _canonical_estimate(estimate)
    if not output_path:
        filename = f"kp_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        output_path = str(DOCS_DIR / filename)

    if company is None:
        company = {"name": "Kolibri Construction"}

    doc_number = _next_doc_number("КП")
    st = _make_styles()
    on_page = _make_page_template("КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ", doc_number, company)

    doc = SimpleDocTemplate(
        output_path, pagesize=A4,
        topMargin=22 * mm, bottomMargin=22 * mm,
        leftMargin=20 * mm, rightMargin=20 * mm,
    )

    elements = []

    # Title block
    elements.append(Paragraph("КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ", st["title"]))
    elements.append(Paragraph(f"№ {doc_number} от {datetime.now().strftime('%d.%m.%Y')}", st["small"]))
    elements.append(Spacer(1, 4 * mm))
    elements.append(HRFlowable(width="100%", thickness=1, color=COL_ACCENT))
    elements.append(Spacer(1, 6 * mm))

    # Company
    elements.extend(_company_header_block(company, st))
    elements.append(Spacer(1, 4 * mm))

    # Client
    client = estimate.get("client", {})
    if client.get("name"):
        elements.append(Paragraph(f"<b>Для:</b> {client['name']}", st["normal"]))
    if client.get("phone"):
        elements.append(Paragraph(f"<b>Телефон:</b> {client['phone']}", st["normal"]))
    elements.append(Spacer(1, 6 * mm))

    # Subject
    elements.append(Paragraph(f"Предмет: {_safe_text(estimate.get('title', 'Строительные работы'))}", st["subtitle"]))
    elements.append(Spacer(1, 4 * mm))

    # Items
    items = estimate.get("items", [])
    totals = estimate.get("totals", {})
    if items:
        elements.extend(_build_items_table(items, totals, st, show_sections=True, sections=estimate.get("sections", [])))
        elements.append(Spacer(1, 8 * mm))

    # Terms
    elements.append(HRFlowable(width="100%", thickness=0.5, color=COL_LIGHT_GREY))
    elements.append(Spacer(1, 4 * mm))
    elements.append(Paragraph("<b>Условия:</b>", st["bold"]))
    elements.append(Paragraph("&bull; Срок выполнения: по согласованию", st["normal"]))
    elements.append(Paragraph("&bull; Гарантия: 12 месяцев", st["normal"]))
    elements.append(Paragraph("&bull; Оплата: поэтапная, по факту выполнения", st["normal"]))

    # Validity
    elements.append(Spacer(1, 4 * mm))
    elements.append(Paragraph("Данное предложение действительно в течение 10 календарных дней.", st["small"]))

    doc.build(elements, onFirstPage=on_page, onLaterPages=on_page)
    return output_path


# ---------------------------------------------------------------------------
# Act of Completed Work (Акт) PDF
# ---------------------------------------------------------------------------

def generate_act_pdf(estimate: dict, company: dict = None, output_path: Optional[str] = None) -> str:
    estimate = _canonical_estimate(estimate)
    if not output_path:
        filename = f"act_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        output_path = str(DOCS_DIR / filename)

    if company is None:
        company = {"name": "Исполнитель"}

    doc_number = _next_doc_number("АКТ")
    st = _make_styles()
    on_page = _make_page_template("АКТ ВЫПОЛНЕННЫХ РАБОТ", doc_number, company)

    doc = SimpleDocTemplate(
        output_path, pagesize=A4,
        topMargin=22 * mm, bottomMargin=22 * mm,
        leftMargin=20 * mm, rightMargin=20 * mm,
    )

    now = datetime.now()
    client = estimate.get("client", {})
    contractor = company.get("name", "Исполнитель")
    customer = client.get("name", "Заказчик")
    city = _extract_city(client.get("address", "")) or _extract_city(company.get("address", "")) or "Москва"

    elements = []

    # Title
    elements.append(Paragraph("АКТ ВЫПОЛНЕННЫХ РАБОТ", st["title"]))
    elements.append(Paragraph(f"№ {doc_number} от {now.strftime('%d.%m.%Y')}", st["small"]))
    elements.append(Spacer(1, 4 * mm))
    elements.append(HRFlowable(width="100%", thickness=1, color=COL_LIGHT_GREY))
    elements.append(Spacer(1, 6 * mm))

    elements.append(Paragraph(f"г. {city} &nbsp;&nbsp;&nbsp; {_ru_date_long(now)}", st["normal"]))
    elements.append(Spacer(1, 4 * mm))

    elements.append(Paragraph(
        f"Мы, нижеподписавшиеся, со стороны Исполнителя <b>{contractor}</b>, "
        f"с одной стороны, и со стороны Заказчика <b>{customer}</b>, с другой стороны, "
        f"составили настоящий акт о том, что следующие работы выполнены и приняты:",
        st["normal"]
    ))
    elements.append(Spacer(1, 6 * mm))

    # Items
    items = estimate.get("items", [])
    totals = estimate.get("totals", {})
    grand = _money(totals.get("grand_total", sum(_money(it.get("total")) for it in items)))

    if items:
        elements.extend(_build_items_table(items, totals, st, show_sections=False))
        elements.append(Spacer(1, 6 * mm))

    # Summary
    elements.append(Spacer(1, 4 * mm))
    elements.append(Paragraph(f"Общая стоимость выполненных работ: <b>{grand:,.2f} руб.</b>", st["bold"]))
    elements.append(Spacer(1, 4 * mm))
    elements.append(Paragraph(
        "Работы выполнены в полном объёме, в установленные сроки и надлежащего качества. "
        "Стороны претензий друг к другу не имеют.",
        st["normal"]
    ))

    # Signatures
    elements.append(Spacer(1, 12 * mm))
    elements.append(HRFlowable(width="100%", thickness=0.5, color=COL_LIGHT_GREY))
    elements.append(Spacer(1, 4 * mm))
    elements.append(Paragraph("<b>Подписи сторон:</b>", st["bold"]))
    elements.append(Spacer(1, 6 * mm))

    sig_data = [
        [Paragraph("<b>Исполнитель:</b>", st["normal"]), Paragraph("", st["normal"]), Paragraph(f"/ {contractor} /", st["normal"])],
        [Paragraph("М.П.", st["small"]), Paragraph("_____________", st["normal"]), Paragraph("", st["normal"])],
        ["", "", ""],
        [Paragraph("<b>Заказчик:</b>", st["normal"]), Paragraph("", st["normal"]), Paragraph(f"/ {customer} /", st["normal"])],
        [Paragraph("М.П.", st["small"]), Paragraph("_____________", st["normal"]), Paragraph("", st["normal"])],
    ]
    sig_table = Table(sig_data, colWidths=[80, 120, 200])
    sig_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), FONT_NAME),
        ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    elements.append(sig_table)

    doc.build(elements, onFirstPage=on_page, onLaterPages=on_page)
    return output_path


# ---------------------------------------------------------------------------
# Invoice (Счёт) PDF
# ---------------------------------------------------------------------------

def generate_invoice_pdf(estimate: dict, company: dict = None, output_path: Optional[str] = None) -> str:
    estimate = _canonical_estimate(estimate)
    if not output_path:
        filename = f"invoice_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        output_path = str(DOCS_DIR / filename)

    if company is None:
        company = {"name": "Kolibri Construction"}

    doc_number = _next_doc_number("СЧТ")
    st = _make_styles()
    on_page = _make_page_template("СЧЁТ НА ОПЛАТУ", doc_number, company)

    doc = SimpleDocTemplate(
        output_path, pagesize=A4,
        topMargin=22 * mm, bottomMargin=22 * mm,
        leftMargin=20 * mm, rightMargin=20 * mm,
    )

    now = datetime.now()
    client = estimate.get("client", {})
    totals = estimate.get("totals", {})
    grand = _money(totals.get("grand_total", 0))

    elements = []

    # Title
    elements.append(Paragraph("СЧЁТ НА ОПЛАТУ", st["title"]))
    elements.append(Paragraph(f"№ {doc_number} от {now.strftime('%d.%m.%Y')}", st["bold"]))
    elements.append(Spacer(1, 4 * mm))
    elements.append(HRFlowable(width="100%", thickness=1, color=COL_ACCENT))
    elements.append(Spacer(1, 6 * mm))

    # Supplier / Customer
    supp_data = [
        [Paragraph("<b>Поставщик:</b>", st["normal"]), Paragraph(company.get("name", ""), st["bold"])],
        [Paragraph("<b>Покупатель:</b>", st["normal"]), Paragraph(client.get("name", ""), st["bold"])],
    ]
    if company.get("inn"):
        supp_data.insert(1, [Paragraph("ИНН:", st["small"]), Paragraph(company["inn"], st["normal"])])
    if client.get("phone"):
        supp_data.append([Paragraph("Тел. покупателя:", st["small"]), Paragraph(client["phone"], st["normal"])])

    supp_table = Table(supp_data, colWidths=[100, 320])
    supp_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), FONT_NAME),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f7fafc")),
        ("BOX", (0, 0), (-1, -1), 0.5, COL_BORDER),
        ("LINEBELOW", (0, 0), (-1, -2), 0.5, COL_LIGHT_GREY),
    ]))
    elements.append(supp_table)
    elements.append(Spacer(1, 6 * mm))

    # Items
    items = estimate.get("items", [])
    if items:
        elements.extend(_build_items_table(items, totals, st, show_sections=False))
        elements.append(Spacer(1, 6 * mm))

    # Payment total
    elements.append(HRFlowable(width="100%", thickness=1, color=COL_ACCENT))
    elements.append(Spacer(1, 4 * mm))
    pay_data = [
        [Paragraph("Итого к оплате:", st["grand_label"]), Paragraph(f"{grand:,.2f} руб.", st["grand_value"])],
    ]
    pay_table = Table(pay_data, colWidths=[250, 170])
    pay_table.setStyle(TableStyle([
        ("ALIGN", (1, 0), (1, 0), "RIGHT"),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#ebf8ff")),
        ("BOX", (0, 0), (-1, -1), 1, COL_ACCENT),
    ]))
    elements.append(pay_table)

    # Bank details
    if company.get("bank") or company.get("account"):
        elements.append(Spacer(1, 6 * mm))
        elements.append(Paragraph("<b>Банковские реквизиты:</b>", st["bold"]))
        if company.get("bank"):
            elements.append(Paragraph(f"Банк: {company['bank']}", st["normal"]))
        if company.get("account"):
            elements.append(Paragraph(f"Р/с: {company['account']}", st["normal"]))

    # Terms
    elements.append(Spacer(1, 6 * mm))
    elements.append(Paragraph("Срок оплаты: в течение 3 банковских дней с момента получения счёта.", st["small"]))
    elements.append(Paragraph("Назначение платежа: оплата по счёту № {} от {}".format(doc_number, now.strftime("%d.%m.%Y")), st["small"]))

    doc.build(elements, onFirstPage=on_page, onLaterPages=on_page)
    return output_path


# ---------------------------------------------------------------------------
# Estimate DOCX
# ---------------------------------------------------------------------------

def generate_estimate_docx(estimate: dict, company: dict = None, output_path: Optional[str] = None, document: dict = None) -> str:
    estimate = _canonical_estimate(estimate)
    doc_number, doc_date, filename = _document_meta(document, "СМТ", "estimate", "docx")
    if not output_path:
        output_path = str(DOCS_DIR / filename)

    doc = Document()

    # Title
    title = doc.add_heading(estimate.get("title", "Смета"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"№ {doc_number} от {doc_date.strftime('%d.%m.%Y')}")

    # Company
    if company and company.get("name"):
        p = doc.add_paragraph()
        run = p.add_run(company["name"])
        run.bold = True
        run.font.size = Pt(12)
        if company.get("inn"):
            doc.add_paragraph(f"ИНН: {company['inn']}")

    doc.add_paragraph()

    # Client
    client = estimate.get("client", {})
    if client.get("name"):
        doc.add_paragraph(f"Заказчик: {client['name']}")
    if client.get("phone"):
        doc.add_paragraph(f"Телефон: {client['phone']}")
    if client.get("address"):
        doc.add_paragraph(f"Адрес: {client['address']}")

    obj = estimate.get("object", {})
    if obj.get("name") or obj.get("type"):
        object_title = obj.get("name") or obj.get("type")
        area = f", площадь: {obj.get('area', 0)} м²" if obj.get("area") else ""
        doc.add_paragraph(f"Объект: {object_title}{area}")
    if obj.get("address"):
        doc.add_paragraph(f"Адрес объекта: {obj['address']}")
    if estimate.get("region"):
        doc.add_paragraph(f"Регион: {estimate['region']}")

    doc.add_paragraph()

    # Items table
    items = estimate.get("items", [])
    if items:
        table = doc.add_table(rows=1, cols=6)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"

        headers = ["№", "Наименование", "Ед.", "Кол-во", "Цена", "Сумма"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.size = Pt(9)

        idx = 0
        grouped = []
        seen = set()
        for section in sorted(estimate.get("sections", []), key=lambda item: item.get("number", 0)):
            section_items = [item for item in items if item.get("section_id") == section.get("id")]
            if section_items:
                seen.update(id(item) for item in section_items)
                grouped.append((section.get("name", "Раздел"), section_items))
        orphans = [item for item in items if id(item) not in seen]
        if orphans:
            grouped.append(("", orphans))
        for section_name, section_items in grouped or [("", items)]:
            if section_name:
                row = table.add_row()
                row.cells[0].text = str(section_name)
                row.cells[0].merge(row.cells[5])
            for item in section_items:
                idx += 1
                row = table.add_row()
                row.cells[0].text = str(idx)
                row.cells[1].text = item.get("name", "")
                row.cells[2].text = item.get("unit", "")
                row.cells[3].text = str(item.get("quantity", 0))
                row.cells[4].text = f"{_money(item.get('unit_price')):,.2f}"
                row.cells[5].text = f"{_money(item.get('total')):,.2f}"

    # Totals
    totals = estimate.get("totals", {})
    doc.add_paragraph()
    if totals.get("works", 0) > 0:
        doc.add_paragraph(f"Работы: {totals['works']:,.2f} руб.")
    if totals.get("materials", 0) > 0:
        doc.add_paragraph(f"Материалы: {totals['materials']:,.2f} руб.")
    if totals.get("services", 0) > 0:
        doc.add_paragraph(f"Услуги: {totals['services']:,.2f} руб.")
    if totals.get("delivery", 0) > 0:
        doc.add_paragraph(f"Доставка: {totals['delivery']:,.2f} руб.")
    if totals.get("discount", 0) > 0:
        doc.add_paragraph(f"Скидка: -{totals['discount']:,.2f} руб.")

    p = doc.add_paragraph()
    run = p.add_run(f"ИТОГО: {totals.get('grand_total', 0):,.2f} руб.")
    run.font.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph()
    doc.add_paragraph(f"Версия: v{estimate.get('version', 1)}")

    doc.save(output_path)
    _record_document(
        output_path,
        kind="estimate",
        title=estimate.get("title", "Смета"),
        doc_number=doc_number,
        doc_date=doc_date,
        estimate=estimate,
    )
    return output_path


def generate_document_pack(estimate: dict, company: dict = None, document: dict = None) -> list[dict[str, str]]:
    """Create the user-facing estimate document pack from one canonical estimate."""
    pack = [
        ("estimate_pdf", generate_estimate_pdf(estimate, company=company, document=document)),
        ("estimate_docx", generate_estimate_docx(estimate, company=company, document=document)),
        ("commercial_offer", generate_commercial_offer_pdf(estimate, company=company)),
        ("act", generate_act_pdf(estimate, company=company)),
        ("invoice", generate_invoice_pdf(estimate, company=company)),
    ]
    return [
        {
            "kind": kind,
            "filename": Path(path).name,
            "path": path,
            "url": f"/api/documents/file/{quote(Path(path).name)}",
        }
        for kind, path in pack
    ]


# ---------------------------------------------------------------------------
# List documents
# ---------------------------------------------------------------------------

def list_documents() -> list:
    index = _load_document_index()
    docs = []
    for f in DOCS_DIR.iterdir():
        if f.suffix in (".pdf", ".docx"):
            stat = f.stat()
            created_at = datetime.fromtimestamp(stat.st_ctime).isoformat()
            meta = index.get(f.name, {}) if isinstance(index.get(f.name), dict) else {}
            doc_type = meta.get("type") or _document_type(f.name)
            title = meta.get("title") or ""
            if not title and doc_type == "estimate":
                title = _estimate_title_from_messages(stat.st_ctime)
            if not title:
                title = _fallback_document_title(f.name)
            docs.append({
                "filename": f.name,
                "name": title,
                "title": title,
                "type": doc_type,
                "kind": doc_type,
                "has_estimate_preview": doc_type == "estimate",
                "url": f"/api/documents/file/{quote(f.name)}",
                "path": str(f),
                "size": stat.st_size,
                "created": created_at,
                "created_at": meta.get("created_at") or created_at,
                "document_number": meta.get("document_number") or _document_number_from_filename(f.name),
                "document_date": meta.get("document_date"),
                "total": meta.get("total"),
                "items_count": meta.get("items_count"),
                "sections_count": meta.get("sections_count"),
            })
    return sorted(docs, key=lambda x: x["created"], reverse=True)
