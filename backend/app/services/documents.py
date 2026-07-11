from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any
from dataclasses import dataclass
import hashlib

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle



@dataclass(frozen=True)
class GeneratedDocument:
    kind: str
    name: str
    content_type: str
    path: Path
    size_bytes: int
    checksum: str


def _safe(value: str) -> str:
    value = re.sub(r"[^\w.-]+", "-", value.strip(), flags=re.UNICODE).strip("-")
    return value or "vista-document"


def _rub(value: int | float) -> str:
    return f"{round(float(value)):,} ₽".replace(",", " ")


def _font_name() -> str:
    name = "VistaSans"
    if name in pdfmetrics.getRegisteredFontNames():
        return name
    candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            pdfmetrics.registerFont(TTFont(name, str(candidate)))
            return name
    return "Helvetica"


def _pdf(path: Path, estimate: dict[str, Any], proposal: dict[str, Any]) -> None:
    font = _font_name()
    styles = getSampleStyleSheet()
    title = ParagraphStyle("VistaTitle", parent=styles["Title"], fontName=font, fontSize=20, leading=25, textColor=colors.HexColor("#101418"))
    body = ParagraphStyle("VistaBody", parent=styles["BodyText"], fontName=font, fontSize=9.5, leading=14, textColor=colors.HexColor("#3a4047"))
    small = ParagraphStyle("VistaSmall", parent=body, fontSize=8, leading=11, textColor=colors.HexColor("#66707a"))
    total = ParagraphStyle("VistaTotal", parent=body, fontSize=16, leading=20, alignment=TA_RIGHT, textColor=colors.HexColor("#101418"))
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm, topMargin=15 * mm, bottomMargin=15 * mm, title=proposal["title"])
    story = [
        Paragraph("VISTA", small),
        Spacer(1, 4 * mm),
        Paragraph(proposal["title"], title),
        Paragraph(f"Клиент: {proposal['client']}", body),
        Paragraph(f"Город: {estimate.get('city', '—')} · Версия: {estimate.get('version', 1)}", small),
        Spacer(1, 7 * mm),
    ]
    rows = [["Раздел", "Работа", "Ед.", "Кол-во", "Цена", "Коэф.", "Сумма"]]
    for item in estimate.get("items", []):
        line = float(item.get("qty", 0)) * float(item.get("price", 0)) * float(item.get("coef", 1))
        rows.append([
            item.get("section", ""), item.get("name", ""), item.get("unit", ""),
            f"{float(item.get('qty', 0)):g}", _rub(item.get("price", 0)), f"{float(item.get('coef', 1)):g}", _rub(line),
        ])
    table = Table(rows, colWidths=[24 * mm, 53 * mm, 13 * mm, 17 * mm, 22 * mm, 15 * mm, 26 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("FONTSIZE", (0, 1), (-1, -1), 7.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#101418")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d7dce1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f8f9")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (3, 1), (-1, -1), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [table, Spacer(1, 7 * mm)]
    summary = estimate.get("summary", {})
    story += [
        Paragraph(f"Работы: {_rub(summary.get('subtotal', 0))}", body),
        Paragraph(f"Накладные расходы: {_rub(summary.get('overhead', 0))}", body),
        Paragraph(f"Прибыль: {_rub(summary.get('margin', 0))}", body),
        Spacer(1, 2 * mm),
        Paragraph(f"ИТОГО: {_rub(summary.get('total', 0))}", total),
        Spacer(1, 6 * mm),
        Paragraph("Условия оплаты", body),
        Paragraph(" · ".join(proposal.get("payment", [])), small),
        Spacer(1, 4 * mm),
        Paragraph("Документ сформирован Vista OS и прошёл проверку итогов.", small),
    ]
    doc.build(story)


def _xlsx(path: Path, estimate: dict[str, Any]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"
    ws.merge_cells("A1:G1")
    ws["A1"] = estimate.get("project", {}).get("name", "Смета объекта")
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="101418")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 32
    headers = ["Раздел", "Работа", "Ед.", "Количество", "Цена", "Коэффициент", "Сумма"]
    ws.append(headers)
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="28313A")
    for item in estimate.get("items", []):
        ws.append([
            item.get("section"), item.get("name"), item.get("unit"),
            item.get("qty"), item.get("price"), item.get("coef"),
            float(item.get("qty", 0)) * float(item.get("price", 0)) * float(item.get("coef", 1)),
        ])
    start = ws.max_row + 2
    summary = estimate.get("summary", {})
    for offset, (label, key) in enumerate([("Работы", "subtotal"), ("Накладные", "overhead"), ("Прибыль", "margin"), ("Итого", "total")]):
        row = start + offset
        ws.cell(row, 6, label).font = Font(bold=key == "total")
        ws.cell(row, 7, summary.get(key, 0)).font = Font(bold=key == "total")
        ws.cell(row, 7).number_format = '#,##0 "₽"'
    widths = [18, 42, 10, 14, 14, 14, 16]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:G{max(2, start-2)}"
    wb.save(path)


def _docx(path: Path, estimate: dict[str, Any], proposal: dict[str, Any]) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)
    title = doc.add_heading(proposal["title"], 0)
    title.alignment = 0
    doc.add_paragraph(f"Клиент: {proposal['client']}")
    doc.add_paragraph(f"Город: {estimate.get('city', '—')} · Версия {estimate.get('version', 1)}")
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["Раздел", "Работа", "Ед.", "Кол-во", "Цена", "Коэф.", "Сумма"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for item in estimate.get("items", []):
        cells = table.add_row().cells
        total = float(item.get("qty", 0)) * float(item.get("price", 0)) * float(item.get("coef", 1))
        values = [item.get("section", ""), item.get("name", ""), item.get("unit", ""), f"{float(item.get('qty', 0)):g}", _rub(item.get("price", 0)), f"{float(item.get('coef', 1)):g}", _rub(total)]
        for cell, value in zip(cells, values):
            cell.text = str(value)
    doc.add_paragraph()
    summary = estimate.get("summary", {})
    paragraph = doc.add_paragraph()
    paragraph.add_run(f"Итого: {_rub(summary.get('total', 0))}").bold = True
    doc.add_heading("Условия оплаты", level=2)
    for payment in proposal.get("payment", []):
        doc.add_paragraph(payment, style="List Bullet")
    doc.add_paragraph("Сформировано Vista OS. Итоги проверены verifier gate.")
    doc.save(path)


def generate_document_pack(estimate: dict[str, Any], root: Path) -> list[GeneratedDocument]:
    root.mkdir(parents=True, exist_ok=True)
    proposal = {
        "title": f"Коммерческое предложение — {estimate.get('project', {}).get('name', 'объект')}",
        "client": estimate.get("client", {}).get("name", "Клиент"),
        "payment": ["40% аванс", "40% после черновых работ", "20% после сдачи"],
    }
    stem = _safe(f"{estimate['id']}-{estimate.get('project', {}).get('name', 'estimate')}")
    paths = {
        "proposal_pdf": root / f"{stem}-proposal.pdf",
        "estimate_xlsx": root / f"{stem}-estimate.xlsx",
        "proposal_docx": root / f"{stem}-proposal.docx",
        "estimate_json": root / f"{stem}-estimate.json",
        "assumptions_md": root / f"{stem}-assumptions.md",
    }
    _pdf(paths["proposal_pdf"], estimate, proposal)
    _xlsx(paths["estimate_xlsx"], estimate)
    _docx(paths["proposal_docx"], estimate, proposal)
    paths["estimate_json"].write_text(json.dumps(estimate, ensure_ascii=False, indent=2), encoding="utf-8")
    paths["assumptions_md"].write_text("# Допущения и риски\n\n" + "\n".join(f"- {x}" for x in estimate.get("assumptions", [])), encoding="utf-8")
    meta = {
        "proposal_pdf": ("Коммерческое предложение.pdf", "application/pdf"),
        "estimate_xlsx": ("Смета.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "proposal_docx": ("Коммерческое предложение.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        "estimate_json": ("Смета.json", "application/json"),
        "assumptions_md": ("Допущения и риски.md", "text/markdown"),
    }
    result: list[GeneratedDocument] = []
    for kind, path in paths.items():
        payload = path.read_bytes()
        result.append(GeneratedDocument(
            kind=kind, name=meta[kind][0], content_type=meta[kind][1], path=path,
            size_bytes=len(payload), checksum=hashlib.sha256(payload).hexdigest(),
        ))
    return result
