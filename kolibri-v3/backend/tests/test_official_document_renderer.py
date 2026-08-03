"""Tests for the official_ru_v1 deterministic document renderers."""

from __future__ import annotations

import io
import json
import zipfile

from openpyxl import load_workbook
from pypdf import PdfReader

from app.official_document_renderer import (
    RENDERER_VERSION,
    render_document_files,
    render_docx,
    render_pdf,
    render_xlsx,
    render_zip,
)
from app.estimate_document_pack import RenderedDocumentFile


def _snapshot() -> dict:
    return {
        "schemaId": "kolibri.estimate_document_snapshot",
        "schemaVersion": "1.0",
        "tenantId": "t_test",
        "projectId": "p_test",
        "documentId": "d_test",
        "estimateVersion": 1,
        "sourceHash": "sha256:" + "ab" * 32,
        "document": {
            "number": "КП-2026-000001",
            "date": "2026-08-02",
            "validUntil": "2026-08-16",
            "status": "preliminary",
            "mode": "preliminary",
            "rendererVersion": RENDERER_VERSION,
        },
        "project": {"title": "Тестовый дом 100 м²", "status": "active", "updatedAt": "2026-08-02T00:00:00+00:00"},
        "object": {
            "name": "Одноэтажный дом",
            "resolvedLocation": "Казань",
            "country": "Россия",
            "region": "Республика Татарстан",
        },
        "customer": {"displayName": "ООО Заказчик"},
        "contractor": {
            "displayName": "ООО Исполнитель",
            "requisites": {
                "legalName": "ООО Исполнитель",
                "taxId": "5001007322",
                "registrationCode": "500101001",
                "bankName": "АО Банк Тест",
                "bankIdentificationCode": "044525225",
                "settlementAccount": "40702810900000000001",
                "correspondentAccount": "30101810400000000225",
                "legalAddress": "Казань, ул. Тестовая, д. 1",
                "basis": "Устав",
            },
        },
        "lines": [
            {
                "position": 1,
                "id": "row_1",
                "section": "Фундамент",
                "kind": "work",
                "description": "Устройство ленточного фундамента",
                "unit": "м³",
                "quantity": "50",
                "unitPrice": "5000.00",
                "lineTotal": "250000.00",
                "confidence": "source_backed",
                "source": {
                    "sourceType": "source_backed",
                    "label": "Прайс-лист 2026",
                    "region": "Республика Татарстан",
                    "observedAt": "2026-07-01",
                    "stale": False,
                    "sourceHash": "sha256:" + "cd" * 32,
                },
            },
            {
                "position": 2,
                "id": "row_2",
                "section": "Стены",
                "kind": "material",
                "description": "Керамический кирпич М150",
                "unit": "шт.",
                "quantity": "10000",
                "unitPrice": "15.50",
                "lineTotal": "155000.00",
                "confidence": "ai_preliminary",
                "source": {
                    "sourceType": "ai_preliminary",
                    "label": "Предварительная цена AI",
                    "region": None,
                    "observedAt": "2026-08-01",
                    "stale": False,
                    "sourceHash": "sha256:" + "ef" * 32,
                },
            },
        ],
        "sections": [
            {"name": "Фундамент", "total": "250000.00"},
            {"name": "Стены", "total": "155000.00"},
        ],
        "commercialTerms": {
            "currency": "RUB",
            "reservePercent": "10.00",
            "taxPercent": "0.00",
            "discount": "0.00",
            "delivery": "0.00",
            "taxMode": "НДС не выделен",
            "paymentSchedule": [
                {"label": "По выпуску документа", "percent": "100.00"},
            ],
            "directTotal": "405000.00",
            "reserve": "40500.00",
            "tax": "0.00",
            "total": "445500.00",
            "totalWords": "Четыреста сорок пять тысяч пятьсот рублей 00 копеек.",
        },
        "calculationConditions": [
            "Цены указаны на условиях самовывоза со склада поставщика.",
            "Стоимость работ актуальна для стандартных грунтовых условий.",
        ],
        "scopeExclusions": [
            "Благоустройство территории",
            "Подключение инженерных сетей",
        ],
    }


def test_render_pdf_produces_valid_portrait_document() -> None:
    snapshot = _snapshot()
    pdf_bytes = render_pdf(snapshot, requested_kinds=["pack"])
    assert pdf_bytes.startswith(b"%PDF-")
    reader = PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) >= 5
    first_page = reader.pages[0]
    assert float(first_page.mediabox.width) < float(first_page.mediabox.height)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "КОММЕРЧЕСКАЯ СМЕТА" in text
    assert "ЛОКАЛЬНЫЙ СМЕТНЫЙ РАСЧЁТ" in text
    assert "РЕСУРСНАЯ ВЕДОМОСТЬ" in text
    assert "КОНЪЮНКТУРНЫЙ АНАЛИЗ" in text
    assert "ПРИЛОЖЕНИЕ" in text
    assert "PRELIMINARY" in text
    assert "445 500,00 руб." in text
    assert "/JavaScript" not in str(reader.trailer)


def test_render_pdf_commercial_offer_single_page() -> None:
    snapshot = _snapshot()
    pdf_bytes = render_pdf(snapshot, requested_kinds=["commercial_offer"])
    reader = PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) == 1
    text = reader.pages[0].extract_text() or ""
    assert "КОММЕРЧЕСКАЯ СМЕТА" in text


def test_render_pdf_invoice_page_with_requisites() -> None:
    snapshot = _snapshot()
    pdf_bytes = render_pdf(snapshot, requested_kinds=["invoice"])
    reader = PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) == 1
    text = reader.pages[0].extract_text() or ""
    assert "СЧЁТ НА ОПЛАТУ" in text
    assert "БИК" in text
    assert "044525225" in text


def test_render_docx_produces_valid_document() -> None:
    snapshot = _snapshot()
    docx_bytes = render_docx(snapshot, requested_kinds=["pack"])
    assert len(docx_bytes) > 500
    from docx import Document
    doc = Document(io.BytesIO(docx_bytes))
    assert doc.core_properties.author == "Kolibri"
    assert doc.core_properties.title == "Тестовый дом 100 м²"
    paragraph_texts = [paragraph.text for paragraph in doc.paragraphs]
    table_texts = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    all_text = "\n".join(paragraph_texts + table_texts)
    assert "КОММЕРЧЕСКАЯ СМЕТА" in all_text
    assert "ЛОКАЛЬНЫЙ СМЕТНЫЙ РАСЧЁТ" in all_text
    assert "Итого к оплате" in all_text


def test_render_xlsx_produces_valid_workbook() -> None:
    snapshot = _snapshot()
    xlsx_bytes = render_xlsx(snapshot, requested_kinds=["pack"])
    assert len(xlsx_bytes) > 500
    workbook = load_workbook(io.BytesIO(xlsx_bytes))
    sheet_names = set(workbook.sheetnames)
    assert "Коммерческая смета" in sheet_names
    assert "Локальный расчёт" in sheet_names
    assert "Ресурсная ведомость" in sheet_names
    assert "Конъюнктурный анализ" in sheet_names
    assert "Условия" in sheet_names
    commercial = workbook["Коммерческая смета"]
    totals_row = [cell.value for cell in commercial[commercial.max_row]]
    assert "Итого к оплате" in totals_row


def test_render_xlsx_invoice_sheet() -> None:
    snapshot = _snapshot()
    xlsx_bytes = render_xlsx(snapshot, requested_kinds=["invoice"])
    workbook = load_workbook(io.BytesIO(xlsx_bytes))
    assert "Счёт" in workbook.sheetnames
    sheet = workbook["Счёт"]
    values = {str(cell.value) for row in sheet.iter_rows(min_row=2) for cell in row if cell.value}
    assert any("Исполнитель" in v for v in values)


def test_render_zip_contains_manifest_and_all_files() -> None:
    snapshot = _snapshot()
    files = (
        RenderedDocumentFile("pdf", "КП-2026-000001-official-pack.pdf", "application/pdf", b"%PDF-test"),
        RenderedDocumentFile("xlsx", "КП-2026-000001-official-pack.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"xlsx-test"),
        RenderedDocumentFile("docx", "КП-2026-000001-official-pack.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"docx-test"),
    )
    zip_bytes = render_zip(snapshot, files)
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as archive:
        names = set(archive.namelist())
        assert "manifest.json" in names
        assert "КП-2026-000001-official-pack.pdf" in names
        assert "КП-2026-000001-official-pack.xlsx" in names
        assert "КП-2026-000001-official-pack.docx" in names
        manifest = json.loads(archive.read("manifest.json"))
        assert manifest["rendererVersion"] == RENDERER_VERSION
        assert manifest["documentNumber"] == "КП-2026-000001"
        assert manifest["sourceHash"] == snapshot["sourceHash"]
        assert len(manifest["files"]) == 3
        assert all(item["artifactHash"].startswith("sha256:") for item in manifest["files"])


def test_render_zip_uses_stable_timestamps() -> None:
    snapshot = _snapshot()
    files = (RenderedDocumentFile("pdf", "test.pdf", "application/pdf", b"content"),)
    first = render_zip(snapshot, files)
    second = render_zip(snapshot, files)
    with zipfile.ZipFile(io.BytesIO(first)) as a, zipfile.ZipFile(io.BytesIO(second)) as b:
        for name in a.namelist():
            assert a.getinfo(name).date_time == b.getinfo(name).date_time


def test_render_document_files_produces_all_four_formats() -> None:
    snapshot = _snapshot()
    files = render_document_files(snapshot, requested_kinds=["pack"])
    kinds = {file.kind for file in files}
    assert kinds == {"pdf", "xlsx", "docx", "zip"}
    for file in files:
        assert file.filename.startswith("КП-2026-000001")
        assert len(file.content) > 0


def test_render_document_files_invoice_single_format() -> None:
    snapshot = _snapshot()
    files = render_document_files(snapshot, requested_kinds=["invoice"])
    kinds = {file.kind for file in files}
    assert "pdf" in kinds
    assert "zip" in kinds


def test_renderer_version_is_stable() -> None:
    assert RENDERER_VERSION == "official_ru_v1"
